"""院別 月間KPIダッシュボード（シングルページ）

起動方法:
    pip install -r requirements.txt
    streamlit run app.py

読み込めるデータ（サイドバーからアップロード）:
  ① 全体の数値 … 「店舗分析データ (月次推移)」CSV（縦に項目名、横に1月〜12月）、または CRM の月間報告書 Excel
  ② メニュー別の数値 … 「商品売上分析」CSV（商品名・カテゴリー・販売数・売上金額）
目標と比べた差額・進捗率を、単月または「〇月〜〇月」の合算で表示する。
読み込んだデータと目標は、そのブラウザのセッション（st.session_state）だけに保存し、他の人とは共有しない。
"""

import calendar
import hashlib
import hmac
import html
import io
import json
import re
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="月間KPIダッシュボード", page_icon="📊", layout="wide")

# ---------------------------------------------------------------------------
# 定数
# ---------------------------------------------------------------------------
SUM = "合計"   # 積み上がる数値（売上・人数）：日割りの基準で進捗を判定
MEAN = "平均"  # 率・単価・平均値：目標値そのものと比較

SUPPORTED_EXTS = ["xls", "xlsx", "csv"]
EXCEL_ENGINES = {".xlsx": "openpyxl", ".xls": "xlrd"}

ALL_CLINICS = "全院合計"
TOTAL_WORDS = {"合計", "小計", "総計", "計"}

# 判定と色（緑＝達成、青＝ペース内、赤＝遅れ）
ST_DONE, ST_ONTRACK, ST_BEHIND, ST_NOTARGET, ST_MISSING = "達成", "順調", "遅れ", "目標未設定", "未検出"
STATUS_STYLE = {
    ST_DONE: ("#2e7d32", "🟢 達成"),
    ST_ONTRACK: ("#1565c0", "🔵 順調"),
    ST_BEHIND: ("#c62828", "🔴 遅れ"),
    ST_NOTARGET: ("#9e9e9e", "⚪ 目標未設定"),
    ST_MISSING: ("#ef6c00", "⚠️ 未検出"),
}
SALES_COLORS = {"保険": "#4e79a7", "混合": "#b07aa1", "自費": "#f28e2b", "物販": "#59a14f"}


# ---------------------------------------------------------------------------
# 抽出する項目の定義
#   lookups は上から順に試し、最初に値が取れたものを使う。
#   ("cross", シート名に含む文字, 基準ラベル, 行ラベル, 列ラベル)  … 行と列が交わるセル
#   ("right", シート名に含む文字, 基準ラベル, ラベル, None)        … ラベルの右にある最初の数値
# ---------------------------------------------------------------------------
def _sales_items() -> list[dict]:
    items = []
    for patient in ["全体", "新患", "再診", "継続"]:
        for col, base in [("合計", "窓口売上合計"), ("保険", "保険売上"), ("混合", "混合売上"), ("自費", "自費売上"), ("物販", "物販売上")]:
            items.append(
                {
                    "name": base if patient == "全体" else f"{base}（{patient}）",
                    "category": "窓口売上" if patient == "全体" else "患者区分別の売上",
                    "unit": "円",
                    "agg": SUM,
                    "lookups": [
                        ("cross", "窓口売上", None, patient, col),
                        ("cross", "全体1", "窓口売上", patient, col),
                    ],
                }
            )
    return items


CATALOG: list[dict] = _sales_items() + [
    # 単価
    {"name": "窓口金単価", "category": "単価・平均", "unit": "円", "agg": MEAN,
     "lookups": [("cross", "窓口売上", None, "全体", "窓口金単価"), ("cross", "全体1", "窓口売上", "全体", "窓口金単価")]},
    {"name": "自費単価", "category": "単価・平均", "unit": "円", "agg": MEAN,
     "lookups": [("cross", "窓口売上", None, "全体", "自費単価"), ("cross", "全体1", "窓口売上", "全体", "自費単価")]},
    # 単価＝報告書の「全体単価」。保存済みデータに無い場合は 窓口売上合計 ÷ 延べ来院数（切り捨て）で計算する（報告書の値と一致）
    {"name": "全体平均単価", "category": "単価・平均", "unit": "円", "agg": MEAN,
     "ratio_of": ("窓口売上合計", "延べ来院数"), "ratio_scale": 1, "ratio_floor": True,
     "lookups": [("cross", "窓口売上", None, "全体", "全体単価"), ("cross", "全体1", "窓口売上", "全体", "全体単価")]},
    # 患者数
    {"name": "延べ来院数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "来院人数", None, "合計", "人数"), ("cross", "全体1", "来院人数", "合計", "人数")]},
    {"name": "純患者数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "純患者数", None, "純患者数", "当月"), ("cross", "全体1", "純患者数(内訳)", "純患者数", "当月")]},
    {"name": "新患数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "純患者数", None, "新患人数", "当月"), ("cross", "全体1", "純患者数(内訳)", "新患人数", "当月")]},
    {"name": "再診数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "純患者数", None, "再診人数", "当月"), ("cross", "全体1", "純患者数(内訳)", "再診人数", "当月")]},
    {"name": "継続患者数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "純患者数", None, "継続人数", "当月"), ("cross", "全体1", "純患者数(内訳)", "継続人数", "当月")]},
    {"name": "午前来院数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "来院人数", None, "午前", "人数"), ("cross", "全体1", "来院人数", "午前", "人数")]},
    {"name": "午後来院数", "category": "患者数", "unit": "人", "agg": SUM,
     "lookups": [("cross", "来院人数", None, "午後", "人数"), ("cross", "全体1", "来院人数", "午後", "人数")]},
    # 離反率の計算に使う内訳（離反者＝前月来院して当月来院しなかった人、前月の純患者数）
    {"name": "離反者数", "category": "患者数", "unit": "人", "agg": MEAN, "combine": SUM,
     "lookups": [("cross", "純患者数", None, "離反者(前月)", "当月"), ("cross", "全体1", "純患者数(内訳)", "離反者(前月)", "当月")]},
    {"name": "前月純患者数", "category": "患者数", "unit": "人", "agg": MEAN, "combine": SUM,
     "lookups": [("cross", "純患者数", None, "純患者数", "前月"), ("cross", "全体1", "純患者数(内訳)", "純患者数", "前月")]},
    # 率（報告書の 0.545 → 54.5%）
    {"name": "初診率（新患1→2回目）", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "pct": True,
     "lookups": [("cross", "初診率", "新患", "比率", "1-2"), ("cross", "全体1", None, "初診率1-2", "%")]},
    {"name": "継続率（新患2→3回目）", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "pct": True,
     "lookups": [("cross", "初診率", "新患", "比率", "2-3"), ("cross", "全体1", None, "継続率2-3", "%")]},
    {"name": "継続率（新患3→4回目）", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "pct": True,
     "lookups": [("cross", "初診率", "新患", "比率", "3-4"), ("cross", "全体1", None, "継続率3-4", "%")]},
    {"name": "再診の戻り率（1→2回目）", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "pct": True,
     "lookups": [("cross", "初診率", "再診", "比率", "1-2")]},
    {"name": "継続患者の戻り率（1→2回目）", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "pct": True,
     "lookups": [("cross", "初診率", "継続", "比率", "1-2")]},
    # 離反率＝離反者数 ÷ 前月の純患者数（報告書の「離反者(前月)」当月の比率と同じ）。低いほど良い
    {"name": "離反率", "category": "初診率・継続率", "unit": "%", "agg": MEAN, "lower_better": True,
     "ratio_of": ("離反者数", "前月純患者数"), "lookups": []},
    # 平均値
    {"name": "平均通院回数", "category": "単価・平均", "unit": "回", "agg": MEAN,
     "lookups": [("right", "全体1", None, "平均通院回数", None)]},
    {"name": "平均滞在時間", "category": "単価・平均", "unit": "分", "agg": MEAN,
     "lookups": [("right", "全体1", None, "平均滞在時間", None)]},
    {"name": "1日平均来院数", "category": "単価・平均", "unit": "人", "agg": MEAN,
     "lookups": [("right", "全体1", None, "全日(平均)", None)]},
]
CATALOG_SIG = hashlib.md5(repr(CATALOG).encode("utf-8")).hexdigest()
CATEGORY_ORDER = ["窓口売上", "患者数", "初診率・継続率", "単価・平均", "患者区分別の売上", "追加項目"]

# ダッシュボード上部に並べる9項目（内部の項目名 → 画面での表示名）
HEADLINE_KPIS = [
    ("窓口売上合計", "合計売上"),
    ("初診率（新患1→2回目）", "初診率"),
    ("純患者数", "カルテ総数"),          # カルテ総数＝報告書の「純患者数」
    ("新患数", "新患数"),
    ("再診数", "再診数"),
    ("離反率", "離反率"),
    ("1日平均来院数", "平均人数（1日）"),
    ("平均通院回数", "平均通院回数"),
    ("全体平均単価", "単価"),            # 単価＝報告書の「全体単価」
]
HEADLINE_ITEMS = [name for name, _ in HEADLINE_KPIS]

# 初期値の目標（仮の値。画面の表で自由に書き換える）
DEFAULT_TARGETS = {
    "窓口売上合計": 3_000_000, "保険売上": 350_000, "自費売上": 1_000_000, "物販売上": 50_000,
    "延べ来院数": 450, "純患者数": 230, "新患数": 30,
    "初診率（新患1→2回目）": 60, "継続率（新患2→3回目）": 50,
    "窓口金単価": 6_500, "平均通院回数": 2.0,
}

CUSTOM_COLUMNS = ["項目名", "単位", "集計", "シート名", "基準ラベル", "行ラベル", "列ラベル", "割合を%に"]


# ---------------------------------------------------------------------------
# セル値のユーティリティ
# ---------------------------------------------------------------------------
def is_blank(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return not v.strip()
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def norm_text(v) -> str:
    """ラベル比較用の正規化（全角半角・空白・コロンの違いを吸収）。"""
    s = unicodedata.normalize("NFKC", str(v))
    return re.sub(r"\s+", "", s).strip(":").lower()


def parse_number(v) -> float | None:
    """セルの値を数値に変換する。「1,234円」「△500」「41.2%」にも対応。数値でなければ None。"""
    if isinstance(v, (bool, np.bool_)):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        return None if pd.isna(v) else float(v)
    if not isinstance(v, str):
        return None
    s = unicodedata.normalize("NFKC", v).strip()
    if not s:
        return None
    neg = s[0] in "△▲-−" and len(s) > 1
    if neg:
        s = s[1:]
    s = re.sub(r"[,\s円¥%人回分件]", "", s)
    if re.fullmatch(r"\d+(\.\d+)?", s):
        return -float(s) if neg else float(s)
    return None


def col_letter(c: int) -> str:
    s, n = "", c + 1
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


# ---------------------------------------------------------------------------
# ファイル読み込みと全セルの走査
# ---------------------------------------------------------------------------
@st.cache_data
def read_grids(raw: bytes, filename: str) -> dict[str, np.ndarray]:
    """全シートを見出し無しの2次元配列として読み込む。"""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        import csv

        for enc in ("utf-8-sig", "cp932"):
            try:
                rows = list(csv.reader(io.StringIO(raw.decode(enc))))
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("CSVの文字コードを判別できませんでした（UTF-8 または Shift_JIS で保存してください）。")
        width = max((len(r) for r in rows), default=0)
        return {"CSV": np.array([r + [None] * (width - len(r)) for r in rows], dtype=object).reshape(len(rows), width)}
    if ext not in EXCEL_ENGINES:
        raise ValueError(f"対応していないファイル形式です（{ext}）。.xls / .xlsx をアップロードしてください。")
    engine = EXCEL_ENGINES[ext]
    try:
        sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None, header=None, dtype=object, engine=engine)
    except ImportError as e:
        raise ValueError(f"Excelの読み込みに必要な「{engine}」がありません。`pip install {engine}` を実行してください。") from e
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"Excelファイルとして読み込めませんでした（{type(e).__name__}）。ファイルが壊れていないか確認してください。") from e
    return {str(name): sdf.to_numpy(dtype=object) for name, sdf in sheets.items()}


@st.cache_data
def grid_cells(raw: bytes, filename: str) -> dict[str, list[tuple[int, int, str]]]:
    """シートごとに、文字が入っているセルの (行, 列, 正規化した文字) 一覧を作る。"""
    out = {}
    for sheet, grid in read_grids(raw, filename).items():
        out[sheet] = [(r, c, norm_text(v)) for (r, c), v in np.ndenumerate(grid) if isinstance(v, str) and v.strip()]
    return out


def find_cells(cells, label: str, row_range: tuple[int, int] | None = None) -> list[tuple[int, int]]:
    """ラベルに一致するセル。完全一致を優先し、無ければ部分一致。"""
    key = norm_text(label)
    pool = [(r, c, n) for r, c, n in cells if row_range is None or row_range[0] <= r <= row_range[1]]
    exact = [(r, c) for r, c, n in pool if n == key]
    return exact or [(r, c) for r, c, n in pool if key in n]


def lookup_cross(grid, cells, anchor, row_label, col_label):
    """行ラベルと列ラベル（行より上にある見出し）が交わるセルの数値を探す。"""
    anchor_pos, row_range = None, None
    if anchor:
        found = find_cells(cells, anchor)
        if not found:
            return None
        anchor_pos = found[0]
        row_range = (anchor_pos[0], anchor_pos[0] + 40)
    rows = find_cells(cells, row_label, row_range)
    headers = find_cells(cells, col_label, row_range)
    best = None
    for r, rc in rows:
        for hr, hc in headers:
            if not (hr < r <= hr + 12 and hc > rc):
                continue
            v = parse_number(grid[r, hc])
            if v is None:
                continue
            closeness = (r - hr) + (hc - rc)
            tie = abs(r - anchor_pos[0]) + abs(rc - anchor_pos[1]) if anchor_pos else r * 1000 + rc
            key = (closeness, tie)
            if best is None or key < best[0]:
                best = (key, v, (r, hc))
    return None if best is None else (best[1], best[2])


def lookup_right(grid, cells, anchor, label):
    """ラベルの右にある最初の数値。基準ラベルがあれば、その下6行以内にあるラベルだけを使う。"""
    candidates = find_cells(cells, label)
    if anchor:
        found = find_cells(cells, anchor)
        if not found:
            return None
        ar, ac = found[0]
        candidates = sorted(
            [(r, c) for r, c in find_cells(cells, label, (ar, ar + 6))],
            key=lambda rc: (abs(rc[0] - ar) + abs(rc[1] - ac)),
        )
    for r, c in candidates:
        for cc in range(c + 1, min(c + 16, grid.shape[1])):
            v = parse_number(grid[r, cc])
            if v is not None:
                return v, (r, cc)
    return None


def sheets_matching(names, hint: str | None) -> list[str]:
    if not hint:
        return list(names)
    key = norm_text(hint)
    return [s for s in names if key in norm_text(s)]


def read_breakdown(grids, cells_by_sheet) -> pd.DataFrame:
    """「施術分類」の一覧（分類・人数・売上）を読む。専用シート → 全体１ の順。"""
    for hint, name_labels in (("施術分類", ["施術分類"]), ("全体1", ["分類"])):
        for sheet in sheets_matching(grids, hint):
            grid, cells = grids[sheet], cells_by_sheet[sheet]
            for name_label in name_labels:
                for r, c in find_cells(cells, name_label):
                    row_cells = [(cc, n) for rr, cc, n in cells if rr == r and cc > c]
                    col_people = next((cc for cc, n in row_cells if n == "人数"), None)
                    col_sales = next((cc for cc, n in row_cells if n.startswith("売上")), None)
                    if col_people is None or col_sales is None:
                        continue
                    rows = []
                    for rr in range(r + 1, min(r + 200, grid.shape[0])):
                        name = grid[rr, c]
                        if is_blank(name):
                            if rows:
                                break
                            continue
                        label = unicodedata.normalize("NFKC", str(name)).strip()
                        if norm_text(label) in TOTAL_WORDS:
                            break
                        if "未表示" in label:
                            label = "その他（未表示分）"  # 全体１では下位の分類がまとめて表示される
                        elif "合計" in label:
                            continue
                        rows.append(
                            {
                                "分類": label,
                                "人数": parse_number(grid[rr, col_people]) or 0.0,
                                "売上": parse_number(grid[rr, col_sales]) or 0.0,
                            }
                        )
                    if rows:
                        return pd.DataFrame(rows)
    return pd.DataFrame(columns=["分類", "人数", "売上"])


# ---------------------------------------------------------------------------
# 対象期間・対象院の検出
# ---------------------------------------------------------------------------
_D = r"(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})\s*日?\s*(?:\([^)]*\))?"
RANGE_RE = re.compile(_D + r"\s*(?:~|〜|～|－|-|ー|―|から)\s*(?:(\d{4})\s*[/\-.年]\s*)?(\d{1,2})\s*[/\-.月]\s*(\d{1,2})\s*日?")
MONTH_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月(?!\s*\d)")


def safe_date(y, m, d) -> date | None:
    try:
        return date(int(y), int(m), int(d))
    except (TypeError, ValueError):
        return None


def month_bounds(y: int, m: int) -> tuple[date, date] | None:
    if not 1 <= int(m) <= 12:
        return None
    return date(int(y), int(m), 1), date(int(y), int(m), calendar.monthrange(int(y), int(m))[1])


@st.cache_data
def detect_period(raw: bytes, filename: str) -> dict | None:
    """「対象期間 2026/08/01 ～ 2026/08/31」→「2026 年 08 月度」（表紙）→「2026年8月」の順に探す。"""
    grids = read_grids(raw, filename)
    for sheet, grid in grids.items():
        for (r, c), v in np.ndenumerate(grid):
            if isinstance(v, str) and (m := RANGE_RE.search(unicodedata.normalize("NFKC", v))):
                y1, m1, d1, y2, m2, d2 = m.groups()
                s, e = safe_date(y1, m1, d1), safe_date(y2 or y1, m2, d2)
                if s and e:
                    return {"start": min(s, e), "end": max(s, e), "where": f"{sheet}!{col_letter(c)}{r + 1}"}
    for sheet, grid in grids.items():  # 表紙：「2026」「年」「08」「月度」が別々のセル
        for (r, c), v in np.ndenumerate(grid):
            if isinstance(v, str) and norm_text(v) == "年":
                left = [parse_number(grid[r, cc]) for cc in range(max(0, c - 12), c)]
                right = [parse_number(grid[r, cc]) for cc in range(c + 1, min(c + 12, grid.shape[1]))]
                y = next((x for x in reversed(left) if x and 2000 <= x <= 2100), None)
                mo = next((x for x in right if x and 1 <= x <= 12), None)
                if y and mo and (b := month_bounds(int(y), int(mo))):
                    return {"start": b[0], "end": b[1], "where": f"{sheet}!{col_letter(c)}{r + 1}"}
    for sheet, grid in grids.items():
        for (r, c), v in np.ndenumerate(grid):
            if isinstance(v, str) and (m := MONTH_RE.search(unicodedata.normalize("NFKC", v))):
                if b := month_bounds(*m.groups()):
                    return {"start": b[0], "end": b[1], "where": f"{sheet}!{col_letter(c)}{r + 1}"}
    return None


def detect_clinic(grids, cells_by_sheet, filename: str) -> str:
    for sheet, cells in cells_by_sheet.items():
        for r, c in find_cells(cells, "対象院"):
            grid = grids[sheet]
            for cc in range(c + 1, min(c + 15, grid.shape[1])):
                v = grid[r, cc]
                if isinstance(v, str) and v.strip() and "${" not in v:
                    return v.strip()
    return Path(filename).stem


# ---------------------------------------------------------------------------
# 1ファイル分の抽出
# ---------------------------------------------------------------------------
def ratio_value(spec: dict, numerator: float, denominator: float) -> float:
    """計算項目の値（離反率は ×100 の%、単価は割り算して切り捨て）。"""
    value = numerator / denominator * spec.get("ratio_scale", 100)
    return float(int(value)) if spec.get("ratio_floor") else value


def build_catalog(custom: tuple) -> list[dict]:
    catalog = [dict(x) for x in CATALOG]
    for name, unit, agg, sheet, anchor, row, col, pct in custom:
        spec = {
            "name": name,
            "category": "追加項目",
            "unit": unit,
            "agg": agg,
            "pct": bool(pct),
            "lookups": [("cross" if col else "right", sheet or None, anchor or None, row, col or None)],
        }
        existing = next((x for x in catalog if x["name"] == name), None)
        if existing:  # 同名の項目は、追加した探し方を先に試す
            existing["lookups"] = spec["lookups"] + existing["lookups"]
        else:
            catalog.append(spec)
    return catalog


@st.cache_data
def extract_file(raw: bytes, filename: str, custom: tuple, catalog_sig: str) -> dict:
    """catalog_sig は項目定義が変わったときにキャッシュを作り直すためのキー。"""
    grids = read_grids(raw, filename)
    cells = grid_cells(raw, filename)
    values = {}
    for spec in build_catalog(custom):
        got = None
        for kind, hint, anchor, row, col in spec["lookups"]:
            for sheet in sheets_matching(grids, hint):
                if kind == "cross":
                    res = lookup_cross(grids[sheet], cells[sheet], anchor, row, col)
                else:
                    res = lookup_right(grids[sheet], cells[sheet], anchor, row)
                if res:
                    (r, c) = res[1]
                    got = (res[0], f"{sheet}!{col_letter(c)}{r + 1}")
                    break
            if got:
                break
        if got:
            v = got[0] * 100 if spec.get("pct") and abs(got[0]) <= 1.5 else got[0]
            values[spec["name"]] = {"value": v, "found": True, "where": got[1]}
        else:
            values[spec["name"]] = {"value": 0.0, "found": False, "where": ""}
    for spec in build_catalog(custom):  # 他の項目から計算する項目（例：離反率＝離反者数÷前月純患者数）
        if spec.get("ratio_of") and not values[spec["name"]]["found"]:  # 報告書から直接取れた値があればそれを優先
            num, den = (values.get(n, {"found": False}) for n in spec["ratio_of"])
            if num["found"] and den["found"] and den["value"]:
                values[spec["name"]] = {"value": ratio_value(spec, num["value"], den["value"]), "found": True,
                                        "where": f"{num['where']} ÷ {den['where']}"}
    return {
        "values": values,
        "period": detect_period(raw, filename),
        "clinic": detect_clinic(grids, cells, filename),
        "breakdown": read_breakdown(grids, cells),
    }


# ---------------------------------------------------------------------------
# サンプルの月間報告書（実際の帳票と同じシート構成）
# ---------------------------------------------------------------------------
def make_sample_report(start: date, end: date, perf: float, clinic: str = "サンプル整骨院") -> bytes | None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return None
    frac = ((end - start).days + 1) / calendar.monthrange(start.year, start.month)[1]
    k = frac * perf
    sales = {  # 患者区分 × 区分（保険・混合・自費・物販）
        "新患": [int(38_000 * k), int(80_000 * k), int(230_000 * k), int(3_000 * k)],
        "再診": [int(10_000 * k), int(2_000 * k), int(12_000 * k), 0],
        "継続": [int(270_000 * k), int(1_550_000 * k), int(640_000 * k), int(30_000 * k)],
    }
    sales["全体"] = [sum(x) for x in zip(*sales.values())]
    people = {"新患": int(25 * k), "再診": int(12 * k), "継続": int(185 * k)}
    visits_am, visits_pm = int(200 * k), int(230 * k)
    wb = Workbook()
    cover = wb.active
    cover.title = "表紙"
    cover["A14"] = "月間報告書"
    cover["V19"], cover["AF19"], cover["AJ19"], cover["AP19"] = start.year, "年", f"{start.month:02d}", "月度"

    ws = wb.create_sheet("月間報告書【全体１】")
    ws["A1"], ws["AR1"], ws["AY1"] = "月間報告書【全体１】", "対象期間", f"{start:%Y/%m/%d} ～ {end:%Y/%m/%d}"
    ws["AR2"], ws["AY2"] = "対象院", clinic
    ws["B36"], ws["C36"] = "●", "最高値・最低値・平均"
    for row, (label, value) in enumerate(
        [("全日(平均）", round((visits_am + visits_pm) / max(1, (end - start).days + 1), 1)), ("平均通院回数（回）", round(1.8 + 0.2 * perf, 1)),
         ("平均滞在時間（分）", 48)], start=37
    ):
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=8, value=value)

    ws = wb.create_sheet("窓口売上")
    for c, h in enumerate(["保険", "請求金額", "混合", "自費", "物販", "自費単価", "窓口金単価", "全体単価", "合計"], start=3):
        ws.cell(row=3, column=c, value=h)
    for r, patient in enumerate(["新患", "再診", "継続", "全体"], start=4):
        ins, mix, own, goods = sales[patient]
        n = people.get(patient, sum(people.values()))
        total = ins + mix + own + goods
        for c, v in zip(range(2, 12), [patient, ins, 0, mix, own, goods, int(own / max(n, 1)), int(total / max(n * 2, 1)), int(total / max(n * 2, 1)), total]):
            ws.cell(row=r, column=c, value=v)

    ws = wb.create_sheet("初診率・継続率")
    for top, patient, base in ((3, "新患", people["新患"]), (12, "再診", people["再診"]), (21, "継続", people["継続"])):
        ws.cell(row=top, column=2, value=patient)
        rates = [0.55 * perf, 0.45 * perf, 0.3, 0.2] if patient == "新患" else [0.2, 0.5, 0.3, 0.1] if patient == "再診" else [0.6, 0.48, 0.36, 0.11]
        target_n = base
        for i in range(9):
            ws.cell(row=top, column=3 + i, value=f"{i + 1}-{i + 2}")
            rate = rates[i] if i < len(rates) else 0
            back = int(target_n * rate)
            ws.cell(row=top + 1, column=3 + i, value=target_n)
            ws.cell(row=top + 2, column=3 + i, value=back)
            ws.cell(row=top + 5, column=3 + i, value=round(back / target_n, 3) if target_n else 0)
            target_n = back
        for off, label in enumerate(["対象", "戻り", "浮遊者", "完治者", "比率", "全体"], start=1):
            ws.cell(row=top + off, column=2, value=label)

    ws = wb.create_sheet("純患者数(内訳)")
    for c, h in enumerate(["前月", "比率", "当月", "比率", "増減"], start=3):
        ws.cell(row=3, column=c, value=h)
    pure = sum(people.values())
    for r, (label, cur) in enumerate(
        [("純患者数", pure), ("新患人数", people["新患"]), ("再診人数", people["再診"]), ("継続人数", people["継続"])], start=4
    ):
        prev = int(cur * 1.05)
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=prev)
        ws.cell(row=r, column=5, value=cur)
        ws.cell(row=r, column=7, value=cur - prev)

    ws = wb.create_sheet("来院人数")
    for c, h in zip([2, 3, 4, 7, 8, 9], ["午前/午後", "人数", "比率", "性別", "人数", "比率"]):
        ws.cell(row=3, column=c, value=h)
    for r, (label, v, g_label, g) in enumerate(
        [("午前", visits_am, "男性", int(pure * 0.4)), ("午後", visits_pm, "女性", pure - int(pure * 0.4)), ("合計", visits_am + visits_pm, "合計", pure)], start=4
    ):
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=v)
        ws.cell(row=r, column=7, value=g_label)
        ws.cell(row=r, column=8, value=g)

    ws = wb.create_sheet("施術分類")
    for c, h in enumerate(["施術分類", "人数", "比率", "売上(円)", "比率", "単価(円)"], start=2):
        ws.cell(row=2, column=c, value=h)
    menu = [
        ("保険", 320, sales["全体"][0]), ("テーピング", 80, 26_000),
        ("骨盤矯正", 70, 33_000), ("骨盤矯正体験", 5, 9_000), ("（新規）骨盤矯正回数券", 2, 110_000), ("（継続）骨盤矯正回数券", 1, 30_000),
        ("特別診療", 9, 28_000), ("特別診療声掛け", 3, 0), ("特別診療体験", 38, 133_000),
        ("（新規）特別診療回数券", 7, 470_000), ("（継続）特別診療回数券", 4, 330_000),
        ("鍼", 75, 13_000), ("鍼体験", 6, 18_000), ("（新規）鍼回数券", 7, 280_000), ("（継続）鍼回数券", 3, 85_000),
        ("パイオネックス", 110, 60_000),
        ("EMS", 40, 6_000), ("EMS検査", 4, 4_400), ("（新規）EMS回数券", 1, 99_000),
        ("美容鍼", 2, 0), ("美容鍼体験", 2, 9_900), ("（新規）美容鍼回数券", 1, 55_000),
    ]
    total_people = total_sales = 0
    for r, (name, n, s) in enumerate(menu, start=3):
        n, s = int(n * k), int(s * k) if name != "保険" else s
        total_people += n
        total_sales += s
        for c, v in zip(range(2, 8), [name, n, None, s, None, int(s / max(n, 1))]):
            ws.cell(row=r, column=c, value=v)
    ws.cell(row=3 + len(menu), column=2, value="合計")
    ws.cell(row=3 + len(menu), column=3, value=total_people)
    ws.cell(row=3 + len(menu), column=5, value=total_sales)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@st.cache_data
def sample_files(today: date, version: int = 2) -> list[tuple[str, bytes]]:
    """直近3ヶ月分＋当月途中までのサンプル報告書。"""
    out = []
    first = date(today.year, today.month, 1)
    months = [add_months(first, -3), add_months(first, -2), add_months(first, -1), first]
    for i, start in enumerate(months):
        end = month_bounds(start.year, start.month)[1] if i < 3 else max(start, today - timedelta(days=1))
        data = make_sample_report(start, end, perf=[1.02, 0.95, 1.06, 0.92][i])
        if data:
            out.append((f"月間報告書_サンプル_{start:%Y-%m}.xlsx", data))
    return out


# ---------------------------------------------------------------------------
# データの保存先：このブラウザのセッション（st.session_state）だけ
#   アップロードしたデータ・目標・設定は、他の人の画面とは共有しない。
#   ページを再読み込み（F5）・ログアウトすると消える（目標はファイルに書き出して保存できる）。
# ---------------------------------------------------------------------------
def session_store() -> dict:
    return st.session_state.setdefault("_store", {"targets": {}, "config": {}})


def store_get(key: str) -> dict:
    return dict(session_store().get(key, {}))


def store_put(key: str, value: dict) -> None:
    session_store()[key] = dict(value)


def update_store(key: str, changes: dict) -> None:
    data = store_get(key)
    data.update(changes)
    store_put(key, data)


def save_section(key: str, section: str, value: dict) -> None:
    data = store_get(key)
    data[section] = dict(value)
    store_put(key, data)


# ---------------------------------------------------------------------------
# CSV の読み込み
#   ・店舗分析データ（月次推移）：縦に項目名、横に 1月〜12月
#   ・商品売上分析（メニュー別）：商品名・カテゴリー・販売数・売上金額（期間はファイル名から）
# ---------------------------------------------------------------------------
MONTH_HEADER_RE = re.compile(r"^\s*(\d{1,2})\s*月\s*$")
LABEL_UNIT_RE = re.compile(r"^(.*?)\s*\(([^()]*)\)\s*$")
FILE_PERIOD_RE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})_(\d{4})-(\d{1,2})-(\d{1,2})")

# 9つのKPIなど、ダッシュボードの項目 ← 店舗分析CSVの項目名（単位の（ ）を除いた名前）
CSV_ITEM_MAP = {
    "窓口売上合計": ["合計売上金額"],                 # 「合計売上」のカード
    "初診率（新患1→2回目）": ["新規リピート率"],      # 「初診率」のカード
    "純患者数": ["カルテ総数"],
    "新患数": ["新患数"],
    "再診数": ["再診数"],
    "離反率": ["離反率"],
    "1日平均来院数": ["総来平均"],
    "平均通院回数": ["平均通院日数"],
    "全体平均単価": ["1人あたりの単価", "一人当たりの単価", "一人あたりの単価", "1人当たりの単価"],  # 「単価」のカード
    "延べ来院数": ["総来患数(のべ)", "総来患数"],
    "保険売上": ["保険合計"],
    "自費売上": ["自費合計"],
}
# 上の項目以外でダッシュボードに取り込む区分（支払方法の内訳・経由データなどは数が多いので取り込まない）
CSV_EXTRA_SECTIONS = ("重要指標", "成約データ", "来患データ", "来患分析", "売上分類", "平均単価")
CSV_CATEGORY = "店舗分析CSV"


def decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp932"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSVの文字コードを判別できませんでした（UTF-8 または Shift_JIS で保存してください）。")


def csv_unit(unit: str) -> str:
    return {"円": "円", "人": "人", "枚": "枚", "%": "%", "％": "%", "日": "日", "分": "分", "件": "件", "回": "回"}.get(unit.strip(), unit.strip())


@st.cache_data(show_spinner=False)
def parse_monthly_csv(raw: bytes, filename: str) -> dict | None:
    """店舗分析データ（月次推移）を読む。月の列が3つ以上ある CSV でなければ None。"""
    import csv

    rows = list(csv.reader(io.StringIO(decode_text(raw))))
    if not rows:
        return None
    header = [h.strip() for h in rows[0]]
    month_cols = {i: int(m.group(1)) for i, h in enumerate(header) if (m := MONTH_HEADER_RE.match(h))}
    if len(month_cols) < 3:
        return None
    year_m = next((re.search(r"(\d{4})\s*年", h) for h in header if re.search(r"(\d{4})\s*年", h)), None) or re.search(
        r"(20\d{2})", filename
    )
    year = int(year_m.group(1)) if year_m else date.today().year
    items, section = [], ""
    for r in rows[1:]:
        if len(r) < 2:
            continue
        if r[0].strip():
            section = r[0].strip()
        label = r[1].strip()
        if not label:
            continue
        m = LABEL_UNIT_RE.match(label)
        name, unit = (m.group(1).strip(), csv_unit(m.group(2))) if m else (label, "")
        values = {mo: parse_number(r[i]) for i, mo in month_cols.items() if i < len(r)}
        items.append({"section": section, "name": name, "unit": unit, "values": values})
    return {"year": year, "items": items}


@st.cache_data(show_spinner=False)
def parse_menu_csv(raw: bytes, filename: str) -> dict | None:
    """メニュー別の CSV（商品名・販売数・売上金額など）を読む。該当の列が無ければ None。"""
    df = pd.read_csv(io.StringIO(decode_text(raw)), dtype=str).fillna("")
    cols = {norm_text(c): c for c in df.columns}

    def col(*names):
        return next((cols[norm_text(n)] for n in names if norm_text(n) in cols), None)

    name_col = col("商品名", "メニュー名", "メニュー", "施術分類", "分類")
    count_col = col("販売数", "件数", "人数", "数量")
    sales_col = col("売上金額", "売上", "金額", "売上(円)")
    group_col = col("カテゴリー", "カテゴリ", "グループ", "親メニュー")
    store_col = col("店舗名", "院名", "店舗")
    if not name_col or not (count_col or sales_col):
        return None
    out = pd.DataFrame(
        {
            "分類": df[name_col].map(lambda v: unicodedata.normalize("NFKC", str(v)).strip()),
            "グループ": df[group_col].map(lambda v: unicodedata.normalize("NFKC", str(v)).strip() or None) if group_col else None,
            "人数": df[count_col].map(lambda v: parse_number(v) or 0.0) if count_col else 0.0,
            "売上": df[sales_col].map(lambda v: parse_number(v) or 0.0) if sales_col else 0.0,
        }
    )
    out = out[out["分類"] != ""].groupby("分類", as_index=False).agg({"グループ": "first", "人数": "sum", "売上": "sum"})
    period = None
    if m := FILE_PERIOD_RE.search(filename):
        y1, m1, d1, y2, m2, d2 = m.groups()
        start, end = safe_date(y1, m1, d1), safe_date(y2, m2, d2)
        if start and end:
            period = {"start": start, "end": end}
    store = next((v for v in df[store_col] if str(v).strip()), None) if store_col else None
    return {"breakdown": out, "period": period, "store": store}


def month_period(month: pd.Period, today: date) -> dict:
    """月の対象期間。当月は昨日まで（データの最終日）とみなす。"""
    start, end = month_bounds(month.year, month.month)
    if start <= today <= end:
        end = max(start, today - timedelta(days=1))
    return {"start": start, "end": end, "where": "月の列"}


def monthly_csv_records(parsed: dict, filename: str, clinic: str, today: date) -> tuple[list[dict], list[dict]]:
    """店舗分析CSV → 月ごとの実績（records）と、ダッシュボードに追加する項目（specs）。"""
    first_by_name: dict[str, dict] = {}
    for it in parsed["items"]:
        first_by_name.setdefault(norm_text(it["name"]), it)  # 全角半角・空白の違いを無視して照合する
    mapped_labels = {norm_text(label) for labels in CSV_ITEM_MAP.values() for label in labels}
    extra = [it for key, it in first_by_name.items() if it["section"] in CSV_EXTRA_SECTIONS and key not in mapped_labels]
    specs = [
        {
            "name": it["name"],
            "category": CSV_CATEGORY,
            "unit": it["unit"],
            "agg": MEAN if it["unit"] == "%" or any(w in it["name"] for w in ("平均", "率", "単価", "日数")) else SUM,
            "lookups": [],
        }
        for it in extra
    ]
    records = []
    for mo in range(1, 13):
        month = pd.Period(f"{parsed['year']}-{mo:02d}", "M")
        values = {}
        for key, labels in CSV_ITEM_MAP.items():
            it = next((first_by_name[norm_text(l)] for l in labels if norm_text(l) in first_by_name), None)
            v = it["values"].get(mo) if it else None
            values[key] = {"value": float(v) if v is not None else 0.0, "found": v is not None,
                           "where": f"{it['name']}（{mo}月）" if it and v is not None else ""}
        for it in extra:
            v = it["values"].get(mo)
            values[it["name"]] = {"value": float(v) if v is not None else 0.0, "found": v is not None,
                                  "where": f"{it['name']}（{mo}月）" if v is not None else ""}
        if not values["窓口売上合計"]["found"]:  # 合計売上が空の月（まだ来ていない月）は読み込まない
            continue
        records.append({"file": filename, "clinic": clinic, "period": month_period(month, today), "values": values,
                        "breakdown": pd.DataFrame(columns=["分類", "グループ", "人数", "売上"])})
    return records, specs


# ---------------------------------------------------------------------------
# 計算
# ---------------------------------------------------------------------------
def add_months(d: date, months: int) -> date:
    m = d.month - 1 + months
    return date(d.year + m // 12, m % 12 + 1, 1)


def judge(actual: float, target, found: bool, agg: str, ratio: float, lower_better: bool = False) -> dict:
    """差額＝実績−目標、進捗率＝実績÷目標×100。積み上げ項目は日割りの基準と比べて順調／遅れを判定。
    lower_better（離反率など低いほど良い項目）は、目標以下なら達成。進捗率は 目標÷実績×100（下げるほど100%に近づく）。"""
    has_target = target is not None and not pd.isna(target) and float(target) > 0
    out = {"diff": np.nan, "progress": np.nan, "expected": np.nan, "forecast": np.nan}
    if lower_better:
        if has_target:
            target = float(target)
            out["diff"] = actual - target
            out["progress"] = target / actual * 100 if actual > 0 else 100.0
            out["expected"] = target
        out["status"] = ST_MISSING if not found else ST_NOTARGET if not has_target else ST_DONE if actual <= target else ST_BEHIND
        return out
    if has_target:
        target = float(target)
        out["diff"] = actual - target
        out["progress"] = actual / target * 100
        out["expected"] = target * ratio if agg == SUM else target
    if agg == SUM and ratio > 0:
        out["forecast"] = actual / ratio
    if not found:
        status = ST_MISSING
    elif not has_target:
        status = ST_NOTARGET
    elif actual >= target:
        status = ST_DONE
    elif agg == SUM and actual >= out["expected"]:
        status = ST_ONTRACK
    else:
        status = ST_BEHIND
    out["status"] = status
    return out


def fmt(v, unit: str = "") -> str:
    if v is None or pd.isna(v):
        return "—"
    if unit == "%":
        return f"{v:,.1f}%"
    if unit in ("回", "分") or (abs(v) < 100 and not float(v).is_integer()):
        return f"{v:,.1f}{unit}"
    return f"{v:,.0f}{unit}"


def fmt_signed(v, unit: str = "") -> str:
    if v is None or pd.isna(v):
        return "—"
    return ("+" if v >= 0 else "−") + fmt(abs(v), unit)


# ---------------------------------------------------------------------------
# 表示部品
# ---------------------------------------------------------------------------
CSS = """
<style>
.block-container {padding-top:3.2rem; max-width:1240px;}
.page-title {font-size:1.35rem; font-weight:700; margin:0;}
.page-meta {font-size:.82rem; opacity:.62; margin:4px 0 18px;}
.sec {font-size:1rem; font-weight:700; margin:34px 0 12px; padding-bottom:6px;
  border-bottom:1px solid rgba(128,128,128,.2);}
.sec small {font-weight:400; opacity:.6; margin-left:8px; font-size:.78rem;}
.chips {display:flex; gap:8px; flex-wrap:wrap; margin:0 0 16px;}
.chip {padding:5px 12px; border-radius:999px; font-size:.8rem; background:rgba(128,128,128,.08);}
.chip b {font-size:.95rem; margin-left:4px;}
.kpi-grid {display:grid; grid-template-columns:repeat(auto-fit, minmax(210px, 1fr)); gap:12px; margin-bottom:12px;}
.kpi-card {border:1px solid rgba(128,128,128,.18); border-radius:10px; padding:12px 14px 10px;}
.kpi-head {display:flex; align-items:center; gap:6px; font-size:.78rem; opacity:.7; font-weight:600;}
.dot {width:8px; height:8px; border-radius:50%; display:inline-block; flex:none;}
.kpi-value {font-size:1.4rem; font-weight:700; line-height:1.3; margin-top:4px;}
.kpi-card.small .kpi-value {font-size:1.05rem;}
.kpi-alt {font-size:.72rem; opacity:.65; margin-top:-2px;}
.kpi-grid.kpi9 {grid-template-columns:repeat(5, minmax(0, 1fr)); gap:10px;}
.kpi-grid.kpi9 .kpi-card {padding:10px 12px 8px;}
.kpi-grid.kpi9 .kpi-value {font-size:1.2rem;}
@media (max-width: 1100px) {.kpi-grid.kpi9 {grid-template-columns:repeat(3, minmax(0, 1fr));}}
@media (max-width: 640px) {.kpi-grid.kpi9 {grid-template-columns:repeat(2, minmax(0, 1fr));}}
.kpi-gap {font-size:.8rem; font-weight:600; margin-top:2px;}
.kpi-bar {position:relative; height:5px; border-radius:3px; background:rgba(128,128,128,.18); margin:8px 0 6px;}
.kpi-bar > .fill {height:100%; border-radius:3px;}
.kpi-bar > .pace {position:absolute; top:-3px; width:2px; height:11px; background:rgba(90,90,90,.8);}
.kpi-meta {display:flex; justify-content:space-between; font-size:.72rem; opacity:.6;}
.panel {border:1px solid rgba(128,128,128,.18); border-radius:10px; padding:12px 14px;}
.panel h4 {font-size:.82rem; margin:0 0 8px; font-weight:700;}
.panel .row {display:flex; justify-content:space-between; gap:8px; font-size:.82rem; padding:5px 0;
  border-top:1px solid rgba(128,128,128,.12);}
.panel .row:first-of-type {border-top:none;}
.panel .muted {opacity:.6; font-size:.8rem;}
.mt-wrap {overflow-x:auto;}
.mt {min-width:900px; border:1px solid rgba(128,128,128,.18); border-radius:10px; overflow:hidden;}
.mt-row {display:grid; grid-template-columns:minmax(200px,1fr) 64px 76px 128px 110px 108px 132px; gap:10px;
  align-items:center; padding:9px 14px; font-size:.82rem;}
.mt-row > span:nth-child(5) {border-left:1px solid rgba(128,128,128,.18); padding-left:10px;}
.mt-row .gap {display:block; font-size:.68rem; font-weight:600; margin-top:1px;}
.mt-row > span:not(:first-child) {text-align:right;}
.mt-hd {font-size:.72rem; opacity:.6; font-weight:600; background:rgba(128,128,128,.06);}
.mt-group {border-top:1px solid rgba(128,128,128,.14);}
.mt-parent {font-weight:700;}
.mt-parent .amt {font-size:.95rem;}
.mt-child {font-size:.78rem; padding-top:5px; padding-bottom:5px; opacity:.85; background:rgba(128,128,128,.03);}
.mt-child > span:first-child {padding-left:26px;}
.mt-child > span:first-child::before {content:"└ "; opacity:.45;}
details.mt-group > summary {list-style:none; cursor:pointer;}
details.mt-group > summary::-webkit-details-marker {display:none;}
.caret {display:inline-block; width:14px; opacity:.55; transition:transform .15s;}
details[open] .caret {transform:rotate(90deg);}
.cnt {font-weight:400; opacity:.5; font-size:.75rem; margin-left:4px;}
.tsrc {display:block; font-weight:400; opacity:.55; font-size:.68rem;}
.mini {display:inline-flex; align-items:center; gap:6px; justify-content:flex-end; width:100%;}
.mini .track {flex:1; max-width:80px; height:4px; border-radius:2px; background:rgba(128,128,128,.2); overflow:hidden;}
.mini .track > div {height:100%;}
.trial {font-size:.66rem; font-weight:600; margin-left:6px; padding:1px 6px; border-radius:999px; background:rgba(239,108,0,.12); color:#ef6c00;}
</style>
"""


def section(title: str, note: str = "") -> None:
    st.markdown(f"<div class='sec'>{title}{f'<small>{note}</small>' if note else ''}</div>", unsafe_allow_html=True)


def gap_text(row: pd.Series) -> str:
    unit = row["単位"]
    gap = row["目標"] - row["現状"]
    if row.get("低いほど良い"):
        return f"あと {fmt(-gap, unit)} 下げる" if gap < 0 else "達成"
    return f"あと {fmt(gap, unit)}" if gap > 0 else f"達成 {fmt_signed(-gap, row['単位'])}"


def kpi_card(row: pd.Series, ratio: float, small: bool = False, label: str | None = None) -> str:
    color = STATUS_STYLE[row["判定"]][0]
    unit = row["単位"]
    has_target = not pd.isna(row["目標"])
    width = 0 if pd.isna(row["進捗率"]) else max(0.0, min(row["進捗率"], 100.0))
    pace = ""
    if has_target and row["集計"] == SUM and 0 < ratio < 1:
        pace = f"<div class='pace' style='left:{ratio * 100:.1f}%'></div>"
    if not row["取得"]:  # 報告書から取れていない項目（以前に保存した報告書には無い項目など）は判定しない
        has_target, width, pace = False, 0, ""
        gap = "<div class='kpi-gap' style='opacity:.6'>未取得（報告書を再アップロード）</div>"
    elif has_target:
        gap = f"<div class='kpi-gap' style='color:{color}'>{gap_text(row)}</div>"
    else:
        gap = "<div class='kpi-gap' style='opacity:.5'>目標未設定</div>"
    goal = f"目標 {fmt(row['目標'], unit)}{' 以下' if row.get('低いほど良い') else ''}"
    meta = f"<span>{goal}</span><span>{fmt(row['進捗率'], '%')}</span>" if has_target else "<span></span>"
    return (
        f"<div class='kpi-card{' small' if small else ''}'>"
        f"<div class='kpi-head'><span class='dot' style='background:{color}'></span>{label or row['項目']}</div>"
        f"<div class='kpi-value'>{fmt(row['現状'], unit) if row['取得'] else '—'}</div>"
        + (f"<div class='kpi-alt'>{row['別表示']}</div>" if row.get("別表示") else "")
        + f"{gap}"
        f"<div class='kpi-bar'><div class='fill' style='width:{width:.1f}%;background:{color}'></div>{pace}</div>"
        f"<div class='kpi-meta'>{meta}</div></div>"
    )


def list_panel(title: str, rows: pd.DataFrame, color: str, empty: str) -> str:
    if rows.empty:
        body = f"<div class='muted'>{empty}</div>"
    else:
        body = "".join(
            f"<div class='row'><span><span class='dot' style='background:{color}'></span>&nbsp;{r['項目']}</span>"
            f"<span>{gap_text(r)}<span class='muted'>　{fmt(r['進捗率'], '%')}</span></span></div>"
            for _, r in rows.iterrows()
        )
    return f"<div class='panel'><h4>{title}</h4>{body}</div>"


def is_trial_menu(name: str) -> bool:
    """「特別診療体験」「鍼体験」など、体験メニューかどうか。"""
    return "体験" in unicodedata.normalize("NFKC", str(name))


def progress_html(actual, target, unit: str, ratio: float) -> str:
    """進捗バーと進捗率、その下に「あと◯」（達成なら超過分）。目標が無ければ「—」。"""
    if not target:
        return "<span style='opacity:.35'>—</span>"
    res = judge(actual, target, True, SUM, ratio)
    color = STATUS_STYLE[res["status"]][0]
    gap = target - actual
    gap_txt = f"あと {fmt(gap, unit)}" if gap > 0 else ("達成" if gap == 0 else f"達成 +{fmt(-gap, unit)}")
    return (
        f"<span class='mini'><span class='track'><div style='width:{min(res['progress'], 100):.0f}%;background:{color}'></div>"
        f"</span>{res['progress']:.0f}%</span><span class='gap' style='color:{color}'>{gap_txt}</span>"
    )


def menu_tree_html(tree: list[dict], ratio: float, total: float, open_top: int = 3) -> str:
    """親メニュー（合計）の行の下に内訳の行がぶら下がる表。
    件数は内訳ごとに管理するため、親メニュー（合計）の行には件数・件数目標・件数の進捗を出さない。"""
    DASH = "<span style='opacity:.35'>—</span>"

    def dot(row, show_count: bool):
        if row["目標"]:
            status = judge(row["売上"], row["目標"], True, SUM, ratio)["status"]
        elif show_count and row["件数目標"]:
            status = judge(row["人数"], row["件数目標"], True, SUM, ratio)["status"]
        else:
            status = ST_NOTARGET
        return f"<span class='dot' style='background:{STATUS_STYLE[status][0]}'></span>"

    def cells(row, name_html, show_count: bool = True):
        share = row["売上"] / total * 100 if total else 0
        if show_count:
            count_cells = (
                f"<span>{fmt(row['人数'], '件')}</span>"
                f"<span>{fmt(row['件数目標'], '件') if row['件数目標'] else DASH}</span>"
                f"<span>{progress_html(row['人数'], row['件数目標'], '件', ratio)}</span>"
            )
        else:
            count_cells = f"<span>{DASH}</span><span>{DASH}</span><span>{DASH}</span>"
        return (
            f"<span>{name_html}</span>{count_cells}"
            f"<span class='amt' title='構成比 {share:.1f}%'>{fmt(row['売上'], '円')}</span>"
            f"<span>{fmt(row['目標'], '円') if row['目標'] else DASH}</span>"
            f"<span>{progress_html(row['売上'], row['目標'], '円', ratio)}</span>"
        )

    def child_name(m):
        tag = "<span class='trial'>体験</span>" if is_trial_menu(m["name"]) else ""
        return f"{html.escape(m['name'])}{tag}"

    head = (
        "<div class='mt-row mt-hd'><span>メニュー（親）／内訳</span>"
        "<span>件数</span><span>件数目標</span><span>件数の進捗</span>"
        "<span>売上</span><span>売上目標</span><span>売上の進捗</span></div>"
    )
    body = []
    for i, g in enumerate(tree):
        if not g["multi"]:  # 内訳の無い単独メニューは、その分類自体の件数をそのまま表示する
            m = g["members"][0]
            name = f"<span class='caret'></span>{dot(g, True)}&nbsp;{child_name(m)}"
            body.append(f"<div class='mt-group'><div class='mt-row mt-parent'>{cells(g, name)}</div></div>")
            continue
        count = f"<span class='cnt'>{len(g['members'])}項目</span>"
        name = f"<span class='caret'>▸</span>{dot(g, False)}&nbsp;{html.escape(g['group'])} 全体{count}"
        parent = cells(g, name, show_count=False)
        kids = "".join(f"<div class='mt-row mt-child'>{cells(m, child_name(m))}</div>" for m in g["members"])
        body.append(f"<details class='mt-group'{' open' if i < open_top else ''}><summary class='mt-row mt-parent'>{parent}</summary>{kids}</details>")
    return f"<div class='mt-wrap'><div class='mt'>{head}{''.join(body)}</div></div>"


def trial_cards_html(tree: list[dict], ratio: float) -> str:
    """体験メニューごとの件数の進捗カード。"""
    cards = []
    for g in tree:
        for m in g["members"]:
            if not is_trial_menu(m["name"]):
                continue
            target = m["件数目標"]
            if target:
                res = judge(m["人数"], target, True, SUM, ratio)
                color = STATUS_STYLE[res["status"]][0]
                gap = target - m["人数"]
                gap_txt = f"あと {fmt(gap, '件')}" if gap > 0 else ("達成" if gap == 0 else f"達成 +{fmt(-gap, '件')}")
                width = min(res["progress"], 100)
                pace = f"<div class='pace' style='left:{ratio * 100:.1f}%'></div>" if 0 < ratio < 1 else ""
                meta = f"<span>目標 {fmt(target, '件')}</span><span>{res['progress']:.0f}%</span>"
                gap_html = f"<div class='kpi-gap' style='color:{color}'>{gap_txt}</div>"
            else:
                color, width, pace = STATUS_STYLE[ST_NOTARGET][0], 0, ""
                meta = "<span>件数目標を「目標設定」で入力</span><span></span>"
                gap_html = "<div class='kpi-gap' style='opacity:.5'>目標未設定</div>"
            cards.append(
                f"<div class='kpi-card small'>"
                f"<div class='kpi-head'><span class='dot' style='background:{color}'></span>{html.escape(m['name'])}"
                f"<span class='cnt'>{html.escape(g['group'])}</span></div>"
                f"<div class='kpi-value'>{fmt(m['人数'], '件')}</div>{gap_html}"
                f"<div class='kpi-bar'><div class='fill' style='width:{width:.1f}%;background:{color}'></div>{pace}</div>"
                f"<div class='kpi-meta'>{meta}</div></div>"
            )
    return "<div class='kpi-grid'>" + "".join(cards) + "</div>" if cards else ""


# ===========================================================================
# 共通処理（どのページでも実行）：ファイルの読み込み・抽出・集計
# ===========================================================================
st.markdown(CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# パスワード認証（パスワードは .streamlit/secrets.toml の admin_password から読み込む）
# ---------------------------------------------------------------------------
def configured_password() -> str | None:
    try:
        value = st.secrets.get("admin_password")
    except Exception:  # noqa: BLE001  secrets.toml が無い・書式が不正など
        return None
    return str(value) if value else None


def require_login() -> None:
    """未ログインならログイン画面だけを表示して、以降の処理（ファイル読み込み・ダッシュボード）を止める。"""
    if st.session_state.get("authenticated"):
        return
    st.markdown("<div style='height:8vh'></div>", unsafe_allow_html=True)
    _, center, _ = st.columns([1, 1.2, 1])
    with center:
        st.markdown("<div class='page-title'>🔒 月間KPIダッシュボード</div>", unsafe_allow_html=True)
        st.markdown("<div class='page-meta'>パスワードを入力してログインしてください。</div>", unsafe_allow_html=True)
        password = configured_password()
        if password is None:
            st.error("現在ログインできません。管理者にご連絡ください。")
            st.stop()
        with st.form("login_form", border=True):
            entered = st.text_input("パスワード", type="password", placeholder="パスワード")
            submitted = st.form_submit_button("ログイン", type="primary", width="stretch")
        if submitted:
            if hmac.compare_digest(entered.encode("utf-8"), password.encode("utf-8")):
                st.session_state.authenticated = True
                st.rerun()
            st.error("パスワードが違います。")
    st.stop()


require_login()
with st.sidebar:
    if st.button("ログアウト", icon="🔓", key="logout"):
        st.session_state.clear()
        st.rerun()

today = date.today()
ss = st.session_state
ss.setdefault("target_ver", 0)
ss.setdefault("upload_ver", 0)
ss.setdefault("sources", {})  # 読み込んだファイル（このブラウザだけ）：id → 内容

shared_targets = store_get("targets")
config = store_get("config")
ss.kpi_targets = dict(shared_targets.get("kpi", DEFAULT_TARGETS))     # 一度も保存していなければ初期値
ss.menu_targets = dict(shared_targets.get("menu_sales", {}))           # メニュー別 売上目標
ss.menu_count_targets = dict(shared_targets.get("menu_count", {}))     # メニュー別 件数目標

# 追加の抽出項目（目標設定ページで編集。月間報告書の読み込みに使う）
custom = tuple(
    (
        str(r.get("項目名") or "").strip(), str(r.get("単位") or ""), r.get("集計") if r.get("集計") in (SUM, MEAN) else SUM,
        str(r.get("シート名") or "").strip(), str(r.get("基準ラベル") or "").strip(), str(r.get("行ラベル") or "").strip(),
        str(r.get("列ラベル") or "").strip(), bool(r.get("割合を%に")),
    )
    for r in config.get("custom_items", [])
    if str(r.get("項目名") or "").strip() and str(r.get("行ラベル") or "").strip()
)


def add_source(kind: str, name: str, raw: bytes) -> str:
    """アップロードされたファイルを読み、このブラウザのデータに追加する。結果を文字で返す。"""
    sid = f"{name}:{hashlib.md5(raw).hexdigest()[:10]}"
    if sid in ss.sources:
        return f"「{name}」は読み込み済みです"
    ext = Path(name).suffix.lower()
    if kind == "menu":
        if ext != ".csv":
            raise ValueError("メニュー別の数値は CSV（商品売上分析など）をアップロードしてください。")
        parsed = parse_menu_csv(raw, name)
        if parsed is None:
            raise ValueError("「商品名」と「販売数」または「売上金額」の列が見つかりませんでした。")
        month = pd.Period(parsed["period"]["end"], "M") if parsed["period"] else None
        ss.sources[sid] = {"kind": "menu", "name": name, "data": parsed, "month": month, "month_detected": month is not None}
        return f"「{name}」：メニュー {len(parsed['breakdown'])}件"
    if ext == ".csv" and (parsed := parse_monthly_csv(raw, name)) is not None:
        ss.sources[sid] = {"kind": "monthly", "name": name, "data": parsed}
        return f"「{name}」：{parsed['year']}年の月次推移"
    ext_data = extract_file(raw, name, custom, CATALOG_SIG)  # 月間報告書（Excel）
    month = pd.Period(ext_data["period"]["end"], "M") if ext_data["period"] else None
    ss.sources[sid] = {"kind": "report", "name": name, "data": ext_data, "month": month, "month_detected": month is not None}
    return f"「{name}」：{ext_data['clinic']} の月間報告書"


KIND_LABEL = {"monthly": "📈 月次推移CSV", "report": "📄 月間報告書", "menu": "🍽️ メニュー別CSV"}

UPLOAD_AREAS = [
    {
        "kind": "total", "kinds": ("monthly", "report"), "title": "① 全体の数値",
        "label": "店舗分析データCSV（月次推移）・月間報告書Excel", "types": ["csv", "xls", "xlsx"],
        "help": "「店舗分析データ (月次推移)」のCSV（縦に項目名、横に1月〜12月）や、月間報告書のExcelを読み込みます。",
    },
    {
        "kind": "menu", "kinds": ("menu",), "title": "② メニュー別の数値",
        "label": "商品売上分析CSV", "types": ["csv"],
        "help": "商品名・カテゴリー・販売数・売上金額の列があるCSV。対象期間はファイル名（例：_2026-10-01_2026-10-11）から読み取ります。",
    },
]


def month_choice(sid: str, src: dict) -> None:
    """月間報告書・メニュー別CSVの対象月（ファイルから読み取れないときは選んでもらう）。"""
    choices = [pd.Period(today, "M") - i for i in range(24)]
    current = src["month"] or choices[0]
    if current not in choices:
        choices.append(current)
    src["month"] = st.selectbox(
        "対象月" + ("" if src["month_detected"] else "（ファイルから読み取れないため選択）"),
        choices, index=choices.index(current), format_func=lambda m: f"{m.year}年{m.month:02d}月", key=f"month_{sid}",
    )


with st.sidebar:
    st.markdown("**📂 データの読み込み**")
    st.caption("読み込んだデータと目標は、この画面だけで使われます（他の人には見えません）。再読み込みすると消えます。")
    for area in UPLOAD_AREAS:
        with st.container(border=True):
            st.markdown(f"**{area['title']}**")
            files = st.file_uploader(
                area["label"], type=area["types"], accept_multiple_files=True,
                key=f"up_{area['kind']}_{ss.upload_ver}", help=area["help"],
            )
            if files:
                msgs, errors = [], []
                with st.spinner("読み込んでいます…"):
                    for u in files:
                        try:
                            msgs.append(add_source(area["kind"], u.name, u.getvalue()))
                        except ValueError as e:
                            errors.append(f"「{u.name}」：{e}")
                ss[f"notice_{area['kind']}"], ss[f"errors_{area['kind']}"] = msgs, errors
                ss.upload_ver += 1  # アップロード欄を空に戻す
                st.rerun()
            for msg in ss.pop(f"notice_{area['kind']}", []):
                st.success(msg, icon="✅")
            for msg in ss.pop(f"errors_{area['kind']}", []):
                st.error(msg)

            # この欄で読み込み済みのファイル（ファイルごとに削除できる）
            loaded = [(sid, src) for sid, src in ss.sources.items() if src["kind"] in area["kinds"]]
            if loaded:
                st.markdown(f"<small><b>✅ 読み込み済み（{len(loaded)}件）</b></small>", unsafe_allow_html=True)
                for sid, src in loaded:
                    c1, c2 = st.columns([5, 1], vertical_alignment="center")
                    c1.markdown(f"<small>{KIND_LABEL[src['kind']]}<br>{html.escape(src['name'])}</small>", unsafe_allow_html=True)
                    if c2.button("", icon=":material/close:", key=f"del_{sid}", help="このファイルを削除", type="tertiary"):
                        del ss.sources[sid]
                        st.rerun()
                    if src["kind"] in ("menu", "report"):
                        month_choice(sid, src)
                if st.button("この欄をリセット", icon=":material/delete_sweep:", key=f"reset_{area['kind']}", width="stretch",
                             help="この欄で読み込んだデータをすべて削除します"):
                    for sid, _ in loaded:
                        ss.sources.pop(sid, None)
                    ss.upload_ver += 1
                    st.rerun()
            else:
                st.caption("まだ読み込んでいません")

# ---- 読み込んだファイル → 月ごとの実績 ----
csv_clinic = next((src["data"]["store"] for src in ss.sources.values() if src["kind"] == "menu" and src["data"].get("store")), None) \
    or "店舗分析データ"
reports, csv_specs, read_errors = [], [], []
for sid, src in ss.sources.items():
    if src["kind"] == "monthly":
        recs, specs = monthly_csv_records(src["data"], src["name"], csv_clinic, today)
        reports += recs
        csv_specs += [sp for sp in specs if sp["name"] not in {c["name"] for c in csv_specs}]
    elif src["kind"] == "report" and src["month"] is not None:
        ext_data = src["data"]
        period = ext_data["period"] if src["month_detected"] else month_period(src["month"], today)
        reports.append({"file": src["name"], "clinic": ext_data["clinic"], "period": period,
                        "values": {k: dict(v) for k, v in ext_data["values"].items()}, "breakdown": ext_data["breakdown"]})
for src in ss.sources.values():  # メニュー別CSV：同じ月の実績に結合（無ければメニューだけの月を作る）
    if src["kind"] != "menu" or src["month"] is None:
        continue
    period = src["data"]["period"] if src["month_detected"] else month_period(src["month"], today)
    target = next((r for r in reports if pd.Period(r["period"]["end"], "M") == src["month"]
                   and r["clinic"] in (csv_clinic, src["data"].get("store"))), None) or \
        next((r for r in reports if pd.Period(r["period"]["end"], "M") == src["month"]), None)
    menu_df = src["data"]["breakdown"]
    if target is None:
        reports.append({"file": src["name"], "clinic": src["data"].get("store") or csv_clinic, "period": period,
                        "values": {}, "breakdown": menu_df})
    else:
        frames = [f for f in (target.get("menu_csv"), menu_df) if f is not None]
        target["menu_csv"] = pd.concat(frames).groupby("分類", as_index=False).agg({"グループ": "first", "人数": "sum", "売上": "sum"})
        target["breakdown"] = target["menu_csv"]

using_sample = not reports
if using_sample:
    for name, raw in sample_files(today, 2):
        ext = extract_file(raw, name, custom, CATALOG_SIG)
        reports.append({"file": name, "clinic": ext["clinic"], "period": ext["period"],
                        "values": ext["values"], "breakdown": ext["breakdown"]})

# 同じ院・同じ月の報告書が複数ある場合は、対象期間の最終日が一番新しいもの（月途中の累計なら最新）を使う
latest: dict[tuple, dict] = {}
for rep in reports:
    p = rep["period"]
    rep["month"] = pd.Period(p["end"], "M")
    key = (rep["clinic"], rep["month"])
    span = (p["end"], p["end"] - p["start"])
    if key not in latest or span > (latest[key]["period"]["end"], latest[key]["period"]["end"] - latest[key]["period"]["start"]):
        latest[key] = rep
used = list(latest.values())
skipped = [r["file"] for r in reports if r not in used]

# 離反率（＝離反者 ÷ 前月の純患者数）の元になる値をそろえる（月間報告書のデータ用）
CHURN_ALIASES = ("離反者数", "離反者", "離反者(前月)", "離反者（前月）")
by_clinic_month = {(r["clinic"], r["month"]): r for r in used}
for rep in used:
    vals = rep["values"]
    if not vals.get("離反者数", {}).get("found"):
        alias = next((vals[k] for k in CHURN_ALIASES if vals.get(k, {}).get("found")), None)
        if alias:
            vals["離反者数"] = {**alias, "where": alias.get("where") or "離反者"}
    if not vals.get("前月純患者数", {}).get("found"):
        prev = by_clinic_month.get((rep["clinic"], rep["month"] - 1))
        prev_full_month = prev and prev["period"]["end"] == month_bounds(prev["month"].year, prev["month"].month)[1]
        if prev_full_month and prev["values"].get("純患者数", {}).get("found"):
            vals["前月純患者数"] = {"value": prev["values"]["純患者数"]["value"], "found": True,
                                "where": f"前月（{prev['month'].month}月）の報告書の純患者数"}

for msg in read_errors:
    st.error(msg)
if not used:
    st.error("表示できるデータがありません。サイドバーからCSVまたは月間報告書をアップロードしてください。")
    st.stop()

clinics = sorted({r["clinic"] for r in used})
months = sorted({r["month"] for r in used})
month_labels = [f"{m.year}年{m.month:02d}月" for m in months]
MODE_SINGLE, MODE_RANGE = "単月", "期間を指定"
BASIS_SUM, BASIS_AVG = "合算", "月平均"
clinic_options = ([ALL_CLINICS] if len(clinics) > 1 else []) + clinics


def default_view() -> dict:
    return {"mode": MODE_SINGLE, "month": month_labels[-1],
            "range": (month_labels[max(0, len(month_labels) - 3)], month_labels[-1]),
            "basis": BASIS_SUM, "clinic": clinic_options[0]}


def valid_view(v) -> bool:
    return (isinstance(v, dict) and v.get("month") in month_labels and v.get("clinic") in clinic_options
            and all(x in month_labels for x in v.get("range", ())))


if not valid_view(ss.get("applied_view")):
    ss.applied_view = default_view()
view_applied = ss.applied_view

with st.sidebar:
    with st.container(border=True):
        st.markdown("**🔎 表示する期間**")
        # 入力欄の初期値は「いま画面に反映されている設定」
        for key, value in (("pv_mode", view_applied["mode"]), ("pv_month", view_applied["month"]),
                           ("pv_range", tuple(view_applied["range"])), ("pv_basis", view_applied["basis"]),
                           ("pv_clinic", view_applied["clinic"])):
            if key not in ss or (key == "pv_month" and ss[key] not in month_labels) or (key == "pv_clinic" and ss[key] not in clinic_options) \
                    or (key == "pv_range" and not all(x in month_labels for x in ss[key])):
                ss[key] = value
        mode = st.segmented_control("表示のしかた", [MODE_SINGLE, MODE_RANGE], key="pv_mode", label_visibility="collapsed") or MODE_SINGLE
        if mode == MODE_SINGLE:
            st.selectbox("表示する月", month_labels, key="pv_month")
        elif len(month_labels) == 1:
            st.caption("期間を指定できる月が1つしかありません。")
        else:
            st.select_slider("表示する期間（〇月〜〇月）", options=month_labels, key="pv_range")
            st.radio("期間の数値の見方", [BASIS_SUM, BASIS_AVG], key="pv_basis", horizontal=True,
                     help="合算＝期間の合計（目標も月数ぶん）／月平均＝1ヶ月あたりの平均（目標は月間目標のまま）。カードには、もう一方の値も小さく表示します。")
        st.selectbox("対象院", clinic_options, key="pv_clinic")
        pending = {"mode": mode, "month": ss.pv_month, "range": tuple(ss.pv_range), "basis": ss.pv_basis, "clinic": ss.pv_clinic}
        changed = {k: pending[k] for k in pending} != {k: (tuple(v) if k == "range" else v) for k, v in view_applied.items()}
        if st.button("反映", icon=":material/refresh:", type="primary", width="stretch", disabled=not changed,
                     help="選んだ月・期間でダッシュボードの数値を更新します"):
            ss.applied_view = pending
            st.rerun()
        if changed:
            st.caption("⚠️ まだ反映されていません。「反映」を押すとダッシュボードが更新されます。")

# 反映済みの設定で表示する
sel_clinic = view_applied["clinic"]
if view_applied["mode"] == MODE_SINGLE or len(month_labels) == 1:
    month_label = view_applied["month"]
    sel_months = [months[month_labels.index(month_label)]]
else:
    i0, i1 = sorted(month_labels.index(x) for x in view_applied["range"])
    sel_months = months[i0:i1 + 1]
    month_label = f"{month_labels[i0]}〜{month_labels[i1]}" if i0 != i1 else month_labels[i0]
avg_mode = view_applied["basis"] == BASIS_AVG and len(sel_months) > 1  # 期間の「月平均」で表示する

catalog = build_catalog(custom) + [sp for sp in csv_specs if sp["name"] not in {c["name"] for c in CATALOG}]
spec_by_name = {s["name"]: s for s in catalog}
n_months = len(sel_months)
target_scale = 1 if avg_mode else n_months  # 合算は 月間目標 × 月数、月平均は月間目標のまま（合計の項目のみ）


def reports_for(month: pd.Period) -> list[dict]:
    return [r for r in used if r["month"] == month and (sel_clinic == ALL_CLINICS or r["clinic"] == sel_clinic)]


def aggregate_values(reps: list[dict]) -> dict[str, tuple[float, bool, str]]:
    """院・月をまたいだ集計（合計の項目は足し算、率・平均の項目は取得できたものの平均）。"""
    out = {}
    for spec in catalog:
        vals = [r["values"].get(spec["name"], {"value": 0.0, "found": False, "where": ""}) for r in reps]
        found = [v for v in vals if v["found"]]
        if spec.get("combine", spec["agg"]) == SUM:
            value = float(sum(v["value"] for v in vals))
        else:
            value = float(np.mean([v["value"] for v in found])) if found else 0.0
        where = found[0]["where"] if len(reps) == 1 and found else ("" if not found else f"{len(found)}/{len(reps)}件で取得")
        out[spec["name"]] = (value, bool(found), where)
    for spec in catalog:
        # 割合・単価の項目：複数件は「合計どうしの割合」。値が無い場合も、元になる項目から計算する
        if spec.get("ratio_of") and (len(reps) > 1 or not out[spec["name"]][1]):
            (nv, nf, _), (dv, df_, _) = (out.get(n, (0.0, False, "")) for n in spec["ratio_of"])
            if nf and df_ and dv:
                out[spec["name"]] = (ratio_value(spec, nv, dv), True, out[spec["name"]][2] or "合計から計算")
    return out


def breakdown_for(reps: list[dict]) -> pd.DataFrame:
    frames = [r["breakdown"].reindex(columns=["分類", "グループ", "人数", "売上"]) for r in reps if not r["breakdown"].empty]
    if not frames:
        return pd.DataFrame(columns=["分類", "グループ", "人数", "売上"])
    return pd.concat(frames).groupby("分類", as_index=False).agg({"グループ": "first", "人数": "sum", "売上": "sum"})


month_reps = [r for m in sel_months for r in reports_for(m)]
prev_reps = reports_for(sel_months[0] - 1)
if not month_reps:
    st.warning("選んだ期間・院のデータがありません。サイドバーで表示する期間や院を変えてください。")
    st.stop()
period_start = min(r["period"]["start"] for r in month_reps)
period_end = max(r["period"]["end"] for r in month_reps)
last_month = sel_months[-1]
dim = calendar.monthrange(last_month.year, last_month.month)[1]
last_reps = reports_for(last_month)
last_start = date(last_month.year, last_month.month, 1)
last_end = max((r["period"]["end"] for r in last_reps), default=month_bounds(last_month.year, last_month.month)[1])
covered_days = (last_end - last_start).days + 1
last_ratio = min(max(covered_days / dim, 0.0), 1.0)
ratio = (n_months - 1 + last_ratio) / n_months  # 期間合算では、最後の月の経過分だけを日割りする
sum_values = aggregate_values(month_reps)  # 期間の合算（率・平均の項目は月の平均）


def per_month(vals: dict) -> dict:
    """合計の項目を月数で割った「月平均」。率・平均の項目はそのまま。"""
    out = dict(vals)
    for spec in catalog:
        if spec.get("combine", spec["agg"]) == SUM and spec["name"] in out:
            v, f, w = out[spec["name"]]
            out[spec["name"]] = (v / n_months, f, w)
    for spec in catalog:  # 割合・単価は、月平均どうしで計算しても合算と同じ値になる
        if spec.get("ratio_of") and spec["name"] in vals:
            out[spec["name"]] = vals[spec["name"]]
    return out


avg_values = per_month(sum_values) if n_months > 1 else sum_values
values = avg_values if avg_mode else sum_values
alt_values = (sum_values if avg_mode else avg_values) if n_months > 1 else {}  # カードに小さく出す、もう一方の見方
alt_label = (BASIS_SUM if avg_mode else BASIS_AVG) if n_months > 1 else ""
prev_values = aggregate_values(prev_reps) if prev_reps else {}

# メニュー別：期間を指定しても合算せず、期間の中で最も新しい「メニューのデータがある月」を表示する
menu_month = next((m for m in reversed(sel_months) if any(not r["breakdown"].empty for r in reports_for(m))), sel_months[-1])
menu_now = breakdown_for(reports_for(menu_month))
menu_prev = breakdown_for(reports_for(menu_month - 1))
if menu_month == last_month:
    menu_ratio = last_ratio
else:
    menu_ratio = 1.0

# ---------------------------------------------------------------------------
# 施術分類の階層化（親メニュー → 内訳）
#   「(新規)EMS回数券」「EMS体験」「EMS検査」… → 親メニュー「EMS」
# ---------------------------------------------------------------------------
GROUP_KEY = "group::"  # 親メニューの目標を保存するときのキーの接頭辞
MENU_PREFIX_RE = re.compile(r"^\s*\(\s*(新規|継続|初回|再来|既存)\s*\)\s*")
MENU_SUFFIXES = ["回数券消化カウント用", "消化カウント用", "カウント用", "回数券", "体験", "声掛け", "検査", "チケット"]
# 名前に含まれる文字で親メニューを固定するルール（自動判定より優先。「グループ分けを調整」の手動設定はさらに優先）
FIXED_MENU_GROUPS = [("パイオネックス", "鍼"), ("灸", "鍼")]


def base_menu(name: str) -> str:
    """接頭辞（(新規)など）・接尾辞（回数券・体験・声掛け・検査など）を外したメニューの基本名。"""
    s = unicodedata.normalize("NFKC", str(name)).strip()
    s = MENU_PREFIX_RE.sub("", s)
    s = re.split(r"/", s)[0]
    s = re.sub(r"\s+", "", s)
    for suffix in MENU_SUFFIXES:
        if s.endswith(suffix) and len(s) > len(suffix):
            s = s[: -len(suffix)]
            break
    return s or str(name)


def auto_menu_groups(names) -> dict[str, str]:
    """分類名 → 親メニュー名。基本名が別の親メニュー名で始まる場合（鍼パルス → 鍼）はそちらにまとめる。"""
    bases = {n: base_menu(n) for n in names}
    roots = set(bases.values())
    out = {}
    for n, b in bases.items():
        fixed = next((parent for word, parent in FIXED_MENU_GROUPS if word in unicodedata.normalize("NFKC", n)), None)
        if fixed:
            out[n] = fixed
            continue
        parents = [r for r in roots if r != b and b.startswith(r)]
        out[n] = max(parents, key=len) if parents else b
    return out


ss.menu_groups = dict(config.get("menu_groups", {}))  # 手動で変更したグループ分けだけを保存
all_menu_names = sorted(set(menu_now["分類"]) | set(menu_prev["分類"]))
csv_groups = {
    row["分類"]: row["グループ"] for frame in (menu_prev, menu_now) for _, row in frame.iterrows()
    if isinstance(row.get("グループ"), str) and row.get("グループ")
}  # メニュー別CSVの「カテゴリー」
auto_groups = {**auto_menu_groups(all_menu_names), **csv_groups}
menu_groups = {n: ss.menu_groups.get(n) or auto_groups[n] for n in all_menu_names}


def resolve_group_target(targets: dict, group: str, members: list[dict], field: str, multi: bool):
    """親の目標は「親に直接入れた値」→「内訳の目標の合計」の順に使う。"""
    if not multi:
        return members[0][field]
    explicit = targets.get(GROUP_KEY + group)
    if explicit:
        return explicit
    kids = [m[field] for m in members if m[field]]
    return float(sum(kids)) if kids else None


def menu_tree(scale: int = 1) -> list[dict]:
    """親メニューごとの合計（売上・件数）・目標・内訳。scale は期間合算の月数（月間目標 × 月数）。"""
    sales_targets = {k: v * scale for k, v in ss.menu_targets.items() if v}
    count_targets = {k: v * scale for k, v in ss.menu_count_targets.items() if v}
    now = menu_now.set_index("分類") if not menu_now.empty else pd.DataFrame(columns=["人数", "売上"])
    prev = menu_prev.set_index("分類") if not menu_prev.empty else pd.DataFrame(columns=["人数", "売上"])
    groups: dict[str, list[dict]] = {}
    for name in all_menu_names:
        groups.setdefault(menu_groups[name], []).append(
            {
                "name": name,
                "人数": float(now["人数"].get(name, 0.0)),
                "売上": float(now["売上"].get(name, 0.0)),
                "前月": float(prev["売上"][name]) if name in prev.index else np.nan,
                "前月件数": float(prev["人数"][name]) if name in prev.index else np.nan,
                "目標": sales_targets.get(name),          # 売上目標
                "件数目標": count_targets.get(name),
            }
        )
    tree = []
    for group, members in groups.items():
        members.sort(key=lambda m: (-m["売上"], -m["人数"]))
        multi = len(members) > 1
        prevs = [m["前月"] for m in members if not pd.isna(m["前月"])]
        prev_counts = [m["前月件数"] for m in members if not pd.isna(m["前月件数"])]
        tree.append(
            {
                "group": group,
                "multi": multi,
                "members": members,
                "人数": sum(m["人数"] for m in members),
                "売上": sum(m["売上"] for m in members),
                "前月": sum(prevs) if prevs else np.nan,
                "前月件数": sum(prev_counts) if prev_counts else np.nan,
                "目標": resolve_group_target(sales_targets, group, members, "目標", multi),
                # 件数は内訳ごとに管理する（親メニューの合算はしない）。単独メニューはその分類の件数目標
                "件数目標": None if multi else members[0]["件数目標"],
            }
        )
    tree.sort(key=lambda g: (-g["売上"], -g["人数"], g["group"]))
    return tree


def build_kpi_table() -> pd.DataFrame:
    """目標設定ページの値を参照して、全項目の実績・目標・差額・進捗率・判定を計算する。"""
    rows = []
    for spec in catalog:
        value, found, where = values[spec["name"]]
        target = ss.kpi_targets.get(spec["name"])
        if target and spec["agg"] == SUM:
            target = target * target_scale  # 期間の合算：月間目標 × 月数（月平均は月間目標のまま）
        res = judge(value, target, found, spec["agg"], ratio, spec.get("lower_better", False))
        rows.append(
            {
                "カテゴリ": spec["category"],
                "項目": spec["name"],
                "単位": spec["unit"],
                "集計": spec["agg"],
                "低いほど良い": bool(spec.get("lower_better")),
                "現状": value,
                "取得": found,
                "目標": float(target) if target is not None and not pd.isna(target) else np.nan,
                "差額": res["diff"],
                "進捗率": res["progress"],
                "着地見込み": res["forecast"],
                "判定": res["status"],
                "取得元": where,
                # 期間を指定したとき、合計の項目はもう一方の見方（合算⇔月平均）も表示する
                "別表示": (f"{alt_label} {fmt(alt_values[spec['name']][0], spec['unit'])}"
                         if alt_values and found and spec.get("combine", spec["agg"]) == SUM else ""),
            }
        )
    table = pd.DataFrame(rows)
    order = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    table["_o"] = table["カテゴリ"].map(order).fillna(99)
    return table.sort_values("_o", kind="stable").drop(columns="_o").reset_index(drop=True)


period_text = (
    f"{period_start:%Y/%m/%d}〜{period_end:%m/%d}（{covered_days}/{dim}日・{ratio:.0%}経過）" if n_months == 1
    else f"{sel_months[0].year}年{sel_months[0].month}月〜{sel_months[-1].year}年{sel_months[-1].month}月"
    + (f"（{n_months}ヶ月の月平均・目標は月間目標）" if avg_mode else f"（{n_months}ヶ月の合算・目標も{n_months}ヶ月分）")
)


# ===========================================================================
# ページ：ダッシュボード
# ===========================================================================
def page_dashboard():
    kpi = build_kpi_table()
    by_name = kpi.set_index("項目")
    st.markdown("<div class='page-title'>📊 月間KPIダッシュボード</div>", unsafe_allow_html=True)
    st.markdown(
        f"<div class='page-meta'>{sel_clinic}　·　{period_text}"
        + ("　·　サンプルデータ" if using_sample else "")
        + "</div>",
        unsafe_allow_html=True,
    )
    if using_sample:
        st.info("サンプルを表示中です。サイドバーから店舗分析データのCSVや月間報告書をアップロードしてください。", icon="🧪")

    # ---- 全体の状況（ひと目で分かる要約） ----
    judged = kpi[kpi["判定"].isin([ST_DONE, ST_ONTRACK, ST_BEHIND])]
    counts = {s: int((judged["判定"] == s).sum()) for s in (ST_DONE, ST_ONTRACK, ST_BEHIND)}
    chips = "".join(
        f"<span class='chip'><span class='dot' style='background:{STATUS_STYLE[s][0]}'></span> {s}<b>{counts[s]}</b></span>"
        for s in (ST_DONE, ST_ONTRACK, ST_BEHIND)
    )
    n_missing = int(((kpi["判定"] == ST_MISSING) & kpi["目標"].notna()).sum())  # 目標を設定したのに値が取れていない項目だけ数える
    if n_missing:
        chips += f"<span class='chip' style='opacity:.7'>⚠️ 未検出<b>{n_missing}</b></span>"
    st.markdown(f"<div class='chips'>{chips}</div>", unsafe_allow_html=True)

    cards = [kpi_card(kpi[kpi["項目"] == n].iloc[0], ratio, label=label) for n, label in HEADLINE_KPIS if n in by_name.index]
    st.markdown("<div class='kpi-grid kpi9'>" + "".join(cards) + "</div>", unsafe_allow_html=True)

    behind = judged[judged["判定"] == ST_BEHIND].sort_values("進捗率").head(5)
    good = judged[judged["判定"] == ST_DONE].sort_values("進捗率", ascending=False).head(5)
    c1, c2 = st.columns(2, gap="medium")
    c1.markdown(list_panel("遅れている項目", behind, STATUS_STYLE[ST_BEHIND][0], "遅れている項目はありません"), unsafe_allow_html=True)
    c2.markdown(list_panel("達成した項目", good, STATUS_STYLE[ST_DONE][0], "達成した項目はまだありません"), unsafe_allow_html=True)
    if 0 < ratio < 1:
        st.caption("バーの縦線は、月の経過日数から見て本日時点で到達しているべき位置です。")

    # ---- 項目別の一覧 ----
    section("項目別の目標・実績", "目標は「目標設定」ページで変更できます")
    f1, f2 = st.columns([3, 1])
    with f1:
        scope = st.segmented_control(
            "表示", ["目標を設定した項目", "すべての項目"], default="目標を設定した項目", key="item_scope", label_visibility="collapsed"
        ) or "目標を設定した項目"
    with f2:
        st.page_link(target_page, label="目標を変更する", icon="🎯")
    table = kpi if scope == "すべての項目" else kpi[kpi["目標"].notna()]
    view = pd.DataFrame(
        {
            "判定": table["判定"].map(lambda s: STATUS_STYLE[s][1]),
            "項目": table.apply(lambda r: r["項目"] if r["取得"] else f"⚠️ {r['項目']}", axis=1),
            "実績": table.apply(lambda r: fmt(r["現状"], r["単位"]), axis=1),
            "目標": table.apply(lambda r: fmt(r["目標"], r["単位"]), axis=1),
            "差額": table.apply(lambda r: fmt_signed(r["差額"], r["単位"]), axis=1),
            "進捗率": table["進捗率"],
            "月末着地見込み": table.apply(lambda r: fmt(r["着地見込み"], r["単位"]) if r["集計"] == SUM else "—", axis=1),
        }
    )
    if n_months > 1:  # 期間を指定したときは、もう一方の見方（合算⇔月平均）の列も出す
        view.insert(3, alt_label, table["別表示"].str.replace(f"{alt_label} ", "", regex=False).replace("", "—").values)
    st.dataframe(
        view,
        hide_index=True,
        width="stretch",
        height=min(36 * (len(view) + 1) + 4, 560),
        column_config={
            "判定": st.column_config.TextColumn(width="small"),
            "項目": st.column_config.TextColumn(width="medium"),
            "進捗率": st.column_config.ProgressColumn(format="%.1f%%", min_value=0, max_value=100, color="#4e79a7"),
            "差額": st.column_config.TextColumn(help="実績 − 目標（マイナス＝不足分）"),
        },
    )
    if (~kpi["取得"]).any():
        st.caption("⚠️ の項目は報告書の中に見つからなかったため、0 として計算しています。")

    # ---- 推移 ----
    section("月次推移", "読み込んだすべての月の推移")
    history = []
    for m in months:
        reps = reports_for(m)
        if reps:
            vals = aggregate_values(reps)
            history.append({"月": f"{m.year}-{m.month:02d}", **{k: v[0] for k, v in vals.items()}})
    hist = pd.DataFrame(history)
    if len(hist) < 2:
        st.caption("2ヶ月分以上の月間報告書をアップロードすると、推移グラフが表示されます。")
    else:
        names = [s["name"] for s in catalog]
        t1, t2 = st.columns(2, gap="large")
        with t1:
            item = st.selectbox("項目", names, index=names.index("窓口売上合計"), key="trend_item", label_visibility="collapsed")
            target = ss.kpi_targets.get(item)
            layers = [
                alt.Chart(hist)
                .mark_bar(color="#4e79a7", cornerRadiusEnd=3)
                .encode(
                    x=alt.X("月:N", title=None),
                    y=alt.Y(f"{item}:Q", title=None, axis=alt.Axis(format="~s")),
                    tooltip=["月", alt.Tooltip(f"{item}:Q", format=",.1f")],
                )
            ]
            if target:
                layers.append(alt.Chart(pd.DataFrame({"y": [target]})).mark_rule(color="#2e7d32", strokeDash=[5, 4], size=2).encode(y="y:Q"))
            st.altair_chart(alt.layer(*layers).properties(height=260), width="stretch")
            st.caption(f"{item}（{spec_by_name[item]['unit']}）" + ("・緑の破線＝目標" if target else ""))
        with t2:
            st.markdown("**窓口売上の構成**")
            sales_long = hist.melt(id_vars="月", value_vars=[f"{k}売上" for k in SALES_COLORS if f"{k}売上" in hist.columns],
                                   var_name="区分", value_name="売上")
            sales_long["区分"] = sales_long["区分"].str.replace("売上", "")
            st.altair_chart(
                alt.Chart(sales_long)
                .mark_bar()
                .encode(
                    x=alt.X("月:N", title=None),
                    y=alt.Y("売上:Q", title=None, stack="zero", axis=alt.Axis(format="~s")),
                    color=alt.Color("区分:N", scale=alt.Scale(domain=list(SALES_COLORS), range=list(SALES_COLORS.values())),
                                    legend=alt.Legend(orient="top", title=None)),
                    tooltip=["月", "区分", alt.Tooltip("売上:Q", format=",.0f")],
                )
                .properties(height=232),
                width="stretch",
            )
        if sel_months[-1] == months[-1] and last_ratio < 1:
            st.caption(f"※ {sel_months[-1].month}月は月の途中（{last_ratio:.0%}経過）までの実績です。")

    # ---- メニュー・カテゴリ別 ----
    section("メニュー別", f"{menu_month.year}年{menu_month.month}月のデータ（期間の合算はしません）")
    matrix = [
        {"患者区分": p, "区分": k, "売上": by_name.at[f"{k}売上（{p}）", "現状"]}
        for p in ["新患", "再診", "継続"]
        for k in SALES_COLORS
        if f"{k}売上（{p}）" in by_name.index
    ]
    matrix_df = pd.DataFrame(matrix)
    if not (matrix_df.empty or matrix_df["売上"].sum() == 0):  # 患者区分別の売上（月間報告書のデータがあるときだけ）
        st.markdown("**窓口売上の内訳**")
        left, right = st.columns([5, 6], gap="large")
        left.altair_chart(
            alt.Chart(matrix_df)
            .mark_bar()
            .encode(
                y=alt.Y("患者区分:N", sort=["新患", "再診", "継続"], title=None),
                x=alt.X("売上:Q", title=None, stack="zero", axis=alt.Axis(format="~s")),
                color=alt.Color("区分:N", scale=alt.Scale(domain=list(SALES_COLORS), range=list(SALES_COLORS.values())),
                                legend=alt.Legend(orient="top", title=None)),
                tooltip=["患者区分", "区分", alt.Tooltip("売上:Q", format=",.0f")],
            )
            .properties(height=alt.Step(30)),
            width="stretch",
        )
        pivot = matrix_df.pivot_table(index="患者区分", columns="区分", values="売上", aggfunc="sum").reindex(
            index=["新患", "再診", "継続"], columns=list(SALES_COLORS)
        )
        pivot["合計"] = pivot.sum(axis=1)
        pivot.loc["全体"] = pivot.sum()
        right.dataframe(pivot.style.format("{:,.0f}"), width="stretch")

    st.markdown("**施術分類（メニュー）別の売上**")
    if menu_now.empty:
        st.caption("メニュー別のデータがありません。サイドバーの「② メニュー別の数値」から商品売上分析のCSVを読み込んでください。")
    else:
        tree = menu_tree(scale=1)
        total = sum(g["売上"] for g in tree) or 1.0
        def is_active(g):
            return g["売上"] > 0 or g["人数"] > 0 or bool(g["目標"]) or bool(g["件数目標"])

        active = [g for g in tree if is_active(g)]
        idle = [g for g in tree if not is_active(g)]
        st.markdown(menu_tree_html(active, menu_ratio, total), unsafe_allow_html=True)
        st.caption(
            "親メニューの行は内訳（新規・継続の回数券、体験、声掛けなど）の売上合計で、行をタップすると内訳を開閉できます。"
            "件数は内訳ごとに目標と進捗を表示しています。パイオネックス・お灸は「鍼」に含めています。"
            "目標とグループ分けは「目標設定」ページで変更できます。"
        )
        if idle:
            with st.expander(f"今月の実績が無いメニュー（{len(idle)}件）"):
                st.markdown(menu_tree_html(idle, menu_ratio, total, open_top=0), unsafe_allow_html=True)

        trial_html = trial_cards_html(tree, menu_ratio)
        if trial_html:
            st.markdown("**体験メニューの進捗（件数）**")
            st.markdown(trial_html, unsafe_allow_html=True)
            st.caption("体験メニューごとに、件数目標に対する進捗率と「あと◯件」を表示しています。バーの縦線は本日時点で到達しているべき位置です。")

    with st.expander(f"読み込んだ報告書（{len(reports)}件）"):
        if skipped:
            st.caption(f"同じ院・同じ月の報告書が複数あったため、対象期間が最新のものだけを使っています（不使用：{'、'.join(skipped)}）")
        st.dataframe(
            pd.DataFrame(
                [
                    {
                        "ファイル": r["file"],
                        "対象院": r["clinic"],
                        "対象期間": f"{r['period']['start']:%Y/%m/%d}〜{r['period']['end']:%Y/%m/%d}",
                        "取得できた項目": f"{sum(v['found'] for v in r['values'].values())} / {len(r['values'])}",
                        "使用": "✅" if r in used else "—",
                    }
                    for r in reports
                ]
            ),
            hide_index=True,
            width="stretch",
        )


# ===========================================================================
# ページ：目標設定（目標の入力はここだけ）
# ===========================================================================
def page_targets():
    st.markdown("<div class='page-title'>🎯 目標設定</div>", unsafe_allow_html=True)
    st.markdown(
        "<div class='page-meta'>ここで保存した<b>月間</b>目標は、この画面のダッシュボード（カード・一覧・メニュー別・推移）で使われます"
        "（他の人とは共有されません）。期間を合算して表示するときは、合計の項目の目標が月数ぶんになります。"
        f"参考として {month_label}（{sel_clinic}）の実績を表示しています。</div>",
        unsafe_allow_html=True,
    )
    num = st.column_config.NumberColumn

    kpi_rows = [
        {
            "カテゴリ": spec["category"],
            "項目": spec["name"],
            "単位": spec["unit"],
            "前月実績": prev_values.get(spec["name"], (np.nan,))[0] if prev_values else np.nan,
            "今月実績": values[spec["name"]][0],
            "目標": ss.kpi_targets.get(spec["name"]),
        }
        for spec in catalog
    ]
    kpi_df = pd.DataFrame(kpi_rows)
    kpi_df["目標"] = pd.to_numeric(kpi_df["目標"], errors="coerce")
    order = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    kpi_df = kpi_df.sort_values("カテゴリ", key=lambda s: s.map(order).fillna(99), kind="stable").reset_index(drop=True)

    # 施術分類の目標：親メニュー（合計）の行の下に内訳の行を並べる（件数・売上の両方）
    MENU_COLS = ["_key", "_group", "_前月件数", "_前月売上", "階層", "メニュー", "今月件数", "件数目標", "今月売上", "売上目標"]

    def menu_row(key, group, level, label, row):
        parent = level == "親（合計）"  # 親メニューの件数は扱わない（空欄）
        return {"_key": key, "_group": group, "_前月件数": None if parent else row["前月件数"], "_前月売上": row["前月"],
                "階層": level, "メニュー": label, "今月件数": None if parent else row["人数"],
                "件数目標": None if parent else row["件数目標"], "今月売上": row["売上"], "売上目標": row["目標"]}

    menu_rows = []
    for g in menu_tree():
        if g["multi"]:
            # 親の目標：親に直接入れた値、無ければ内訳の目標の合計
            menu_rows.append(menu_row(GROUP_KEY + g["group"], g["group"], "親（合計）", f"■ {g['group']} 全体", g))
            for m in g["members"]:
                label = f"　　└ {m['name']}" + ("　🔰体験" if is_trial_menu(m["name"]) else "")
                menu_rows.append(menu_row(m["name"], g["group"], "内訳", label, m))
        else:
            m = g["members"][0]
            menu_rows.append(menu_row(m["name"], g["group"], "単独", f"■ {m['name']}", m))
    menu_df = pd.DataFrame(menu_rows, columns=MENU_COLS)
    for c in ("_前月件数", "_前月売上", "今月件数", "件数目標", "今月売上", "売上目標"):
        menu_df[c] = pd.to_numeric(menu_df[c], errors="coerce")

    with st.form("targets_form", border=False):
        section("KPI項目の目標（月間）", "「目標」の列だけ入力できます。空欄の項目は判定しません")
        kpi_edit = st.data_editor(
            kpi_df if prev_values else kpi_df.drop(columns=["前月実績"]),
            key=f"kpi_target_editor_{ss.target_ver}",
            hide_index=True,
            width="stretch",
            height=min(36 * (len(kpi_df) + 1) + 4, 640),
            disabled=["カテゴリ", "項目", "単位", "前月実績", "今月実績"],
            column_config={
                "カテゴリ": st.column_config.TextColumn(width="small"),
                "項目": st.column_config.TextColumn(width="medium"),
                "単位": st.column_config.TextColumn(width="small"),
                "前月実績": num(format="localized"),
                "今月実績": num(format="localized"),
                "目標": num("🎯 目標", format="localized", min_value=0.0),
            },
        )
        section("施術分類（メニュー）別の目標（月間）", "「件数目標」「売上目標」の列に入力します")
        st.markdown(
            "<div class='page-meta' style='margin:-4px 0 10px'>"
            "・<b>内訳に入れる</b> → 親メニューの目標は、内訳の目標の合計に自動でなります（保存すると更新）<br>"
            "・<b>親メニューに入れる</b> → グループ全体の目標になります。内訳の合計より優先されます<br>"
            "・親メニューを<b>空欄にして保存</b>すると、内訳の合計に戻ります<br>"
            "・<b>件数目標は内訳・単独メニューの行に入力します</b>（親メニューの行の件数は使いません）。🔰体験 の行に入れると、ダッシュボードに体験メニューの進捗が表示されます</div>",
            unsafe_allow_html=True,
        )
        menu_edit = st.data_editor(
            menu_df,
            key=f"menu_target_editor_{ss.target_ver}",
            hide_index=True,
            width="stretch",
            height=min(36 * (len(menu_df) + 1) + 4, 640),
            column_order=["階層", "メニュー", "今月件数", "件数目標", "今月売上", "売上目標"],
            disabled=["_key", "_group", "_前月件数", "_前月売上", "階層", "メニュー", "今月件数", "今月売上"],
            column_config={
                "階層": st.column_config.TextColumn(width="small"),
                "メニュー": st.column_config.TextColumn(width="medium"),
                "今月件数": num("今月件数（件）", format="localized"),
                "件数目標": num("🎯 件数目標（件）", format="localized", min_value=0.0, step=1.0),
                "今月売上": num("今月売上（円）", format="localized"),
                "売上目標": num("🎯 売上目標（円）", format="localized", min_value=0.0),
            },
        )
        b1, b2, _ = st.columns([1, 1, 3])
        saved = b1.form_submit_button("💾 目標を保存", type="primary", width="stretch")
        fill_prev = b2.form_submit_button("空欄を前月実績で埋める", width="stretch", disabled=not prev_values,
                                          help="目標が空欄の項目に、前月の実績値を仮の目標として入れます（保存も同時に行います）")

    if saved or fill_prev:
        kpi_edit = kpi_edit.copy()
        menu_edit = menu_edit.copy()
        if fill_prev and "前月実績" in kpi_edit.columns:
            kpi_edit["目標"] = kpi_edit["目標"].fillna(kpi_edit["前月実績"])
        leaf = menu_edit["階層"] != "親（合計）"
        if fill_prev:
            menu_edit.loc[leaf, "件数目標"] = menu_edit.loc[leaf, "件数目標"].fillna(menu_edit.loc[leaf, "_前月件数"])
            menu_edit.loc[leaf, "売上目標"] = menu_edit.loc[leaf, "売上目標"].fillna(menu_edit.loc[leaf, "_前月売上"])
        new_kpi = {r["項目"]: float(r["目標"]) for _, r in kpi_edit.iterrows() if not pd.isna(r["目標"]) and r["目標"] > 0}
        new_kpi.update({k: v for k, v in ss.kpi_targets.items() if k not in set(kpi_edit["項目"])})  # 表に無い項目の目標は残す
        shown_keys = set(menu_edit["_key"])

        def collect(column: str, old: dict, parents: bool = True) -> dict:
            new = {k: v for k, v in old.items() if k not in shown_keys}  # 表示していないメニュー（他院など）の目標は残す
            for _, r in menu_edit[leaf].iterrows():
                if not pd.isna(r[column]) and r[column] > 0:
                    new[r["_key"]] = float(r[column])
            # 親の目標は、内訳の合計と違う値が入っているときだけ「親の目標」として保存する（同じなら内訳の合計に自動で連動）
            for _, r in (menu_edit[~leaf] if parents else menu_edit.iloc[0:0]).iterrows():
                kids_sum = float(menu_edit[leaf & (menu_edit["_group"] == r["_group"])]["_key"].map(new).fillna(0).sum())
                v = r[column]
                if not pd.isna(v) and v > 0 and abs(v - kids_sum) > 0.5:
                    new[r["_key"]] = float(v)
            return new

        new_sales = collect("売上目標", ss.menu_targets)
        new_count = collect("件数目標", ss.menu_count_targets, parents=False)  # 件数は親メニューに保存しない
        # このブラウザのセッションに保存（他の人とは共有しない）
        store_put("targets", {"kpi": new_kpi, "menu_sales": new_sales, "menu_count": new_count})
        saved_now = store_get("targets")
        ss.target_ver += 1
        ss.saved_notice = (
            f"目標を保存しました（KPI {len(saved_now.get('kpi', {}))}項目・"
            f"メニューの売上目標 {len(saved_now.get('menu_sales', {}))}件・件数目標 {len(saved_now.get('menu_count', {}))}件）。"
        )
        st.rerun()
    if notice := ss.pop("saved_notice", None):
        st.success(notice + " ダッシュボードに反映されています。", icon="✅")
        st.page_link(dashboard_page, label="ダッシュボードで確認する", icon="📊")

    with st.expander("🗂 メニューのグループ分けを調整"):
        st.caption(
            "分類名から「(新規)」「(継続)」や「回数券」「体験」「声掛け」「検査」などを外して、自動で親メニューにまとめています"
            "（パイオネックス・お灸は「鍼」に含めます）。"
            "違うグループに入れたい分類は「親メニュー」を書き換えて保存してください（新しい名前を入れると新しいグループになります）。"
        )
        grp_df = pd.DataFrame(
            {"分類": all_menu_names, "自動の親メニュー": [auto_groups[n] for n in all_menu_names], "親メニュー": [menu_groups[n] for n in all_menu_names]},
            dtype="object",  # メニューが0件のときも文字の列として扱う
        )
        with st.form("group_form", border=False):
            grp_edit = st.data_editor(
                grp_df,
                key=f"group_editor_{ss.target_ver}",
                hide_index=True,
                width="stretch",
                height=min(36 * (len(grp_df) + 1) + 4, 480),
                disabled=["分類", "自動の親メニュー"],
                column_config={"親メニュー": st.column_config.TextColumn("✏️ 親メニュー")},
            )
            g1, g2, _ = st.columns([1, 1, 3])
            save_groups = g1.form_submit_button("グループ分けを保存", width="stretch")
            reset_groups = g2.form_submit_button("自動に戻す", width="stretch")
        if save_groups or reset_groups:
            overrides = {k: v for k, v in ss.menu_groups.items() if k not in set(all_menu_names)}
            if save_groups:
                for _, r in grp_edit.iterrows():
                    parent = str(r["親メニュー"] or "").strip()
                    if parent and parent != auto_groups[r["分類"]]:
                        overrides[r["分類"]] = parent
            save_section("config", "menu_groups", overrides)
            ss.target_ver += 1
            st.rerun()

    with st.expander("💾 目標の書き出し・読み込み・リセット"):
        st.caption("目標はこの画面にだけ保存され、再読み込みすると初期値に戻ります。ファイルに書き出しておけば、次回読み込んで続きから使えます。")
        backup = json.dumps({"targets": store_get("targets") or {"kpi": ss.kpi_targets, "menu_sales": ss.menu_targets,
                                                                 "menu_count": ss.menu_count_targets},
                             "config": store_get("config")}, ensure_ascii=False, indent=2)
        e1, e2 = st.columns(2)
        e1.download_button("目標をファイルに書き出す", backup.encode("utf-8"), file_name="kpi_targets.json",
                           mime="application/json", width="stretch")
        if e2.button("目標を初期値に戻す", width="stretch"):
            store_put("targets", {})
            ss.target_ver += 1
            st.rerun()
        loaded = st.file_uploader("書き出した目標ファイル（kpi_targets.json）を読み込む", type=["json"], key=f"targets_file_{ss.target_ver}")
        if loaded is not None:
            try:
                data = json.loads(loaded.getvalue().decode("utf-8"))
                store_put("targets", data.get("targets", {}))
                if data.get("config"):
                    store_put("config", data["config"])
                ss.target_ver += 1
                ss.saved_notice = "目標ファイルを読み込みました。"
                st.rerun()
            except (ValueError, AttributeError):
                st.error("目標ファイルとして読み込めませんでした。")

    with st.expander("⚙️ 抽出項目の追加・調整（上級者向け）"):
        st.caption(
            "標準の項目に無いものを追加できます。「行ラベル」と「列ラベル」が交わるセルの数値を取り出します"
            "（列ラベルが空なら、行ラベルの右にある最初の数値）。シート名は一部だけでも構いません。"
            "追加した項目は、保存した後にアップロードした報告書から取り出されます（すでに保存済みの月は、報告書をアップロードし直してください）。"
        )
        custom_base = pd.DataFrame(config.get("custom_items", []), columns=CUSTOM_COLUMNS)
        custom_base["割合を%に"] = custom_base["割合を%に"].fillna(False).astype(bool)
        with st.form("custom_form", border=False):
            custom_df = st.data_editor(
                custom_base,
                key="custom_editor",
                num_rows="dynamic",
                hide_index=True,
                width="stretch",
                column_config={
                    "項目名": st.column_config.TextColumn(required=True),
                    "単位": st.column_config.TextColumn(default="円"),
                    "集計": st.column_config.SelectboxColumn(options=[SUM, MEAN], default=SUM),
                    "シート名": st.column_config.TextColumn(help="例：窓口売上、全体1"),
                    "基準ラベル": st.column_config.TextColumn(help="同じ名前の行が複数ある場合の目印（例：窓口売上）"),
                    "行ラベル": st.column_config.TextColumn(required=True, help="例：全体"),
                    "列ラベル": st.column_config.TextColumn(help="例：自費"),
                    "割合を%に": st.column_config.CheckboxColumn(default=False, help="0.545 → 54.5% に変換"),
                },
            )
            if st.form_submit_button("追加項目を保存"):
                update_store("config", {"custom_items": custom_df.where(custom_df.notna(), None).to_dict("records")})
                st.rerun()


# ===========================================================================
# ページ切り替え（サイドバー上部）
# ===========================================================================
dashboard_page = st.Page(page_dashboard, title="ダッシュボード", icon="📊", default=True)
target_page = st.Page(page_targets, title="目標設定", icon="🎯", url_path="targets")
st.navigation([dashboard_page, target_page]).run()
